"""Tests for 11_train_random_forest_model_c.py using unittest."""

from __future__ import annotations

import importlib.util
import json
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
from scipy import sparse
from sklearn.model_selection import StratifiedKFold


SCRIPT_PATH = Path(__file__).with_name("11_train_random_forest_model_c.py")
SPEC = importlib.util.spec_from_file_location("train_random_forest_model_c", SCRIPT_PATH)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError(f"Could not load {SCRIPT_PATH}")
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def synthetic_data() -> tuple[sparse.csr_matrix, np.ndarray, pd.DataFrame]:
    """Create deterministic balanced sparse data with class-specific signals."""
    random = np.random.default_rng(42)
    labels = np.repeat(np.asarray(MODULE.CLASS_ORDER, dtype=object), 15)
    values = random.integers(0, 3, size=(45, 12)).astype(float)
    for class_index in range(3):
        rows = slice(class_index * 15, (class_index + 1) * 15)
        values[rows, class_index * 3 : class_index * 3 + 3] += 5
    features = pd.DataFrame(
        {
            "feature_index": np.arange(12),
            "ngram_text": [f"term {index}" for index in range(12)],
            "ngram_length": [2] * 12,
            "idf_value": np.linspace(1.0, 2.0, 12),
        }
    )
    return sparse.csr_matrix(values), labels, features


def synthetic_index(labels: np.ndarray) -> pd.DataFrame:
    """Build a candidate-aligned synthetic test index."""
    return pd.DataFrame(
        {
            "matrix_row_index": np.arange(len(labels)),
            "candidate_id": [f"TEST_{index:03d}" for index in range(len(labels))],
            "candidate_text_original": [f"Synthetic statement {index}" for index in range(len(labels))],
            "reviewer_gold_label": labels,
            "split": "test",
        }
    )


class ModelCTests(unittest.TestCase):
    def test_proportional_minimum_reproduces_registered_calculation(self) -> None:
        self.assertEqual(MODULE.proportional_minimum_feature_count(23_464), 124)
        self.assertEqual(round(100 * 20 / 3_770, 4), 0.5305)
        self.assertEqual(round(100 * 124 / 23_464, 4), 0.5285)

    def test_registry_migration_preserves_rows_and_allows_eight_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as temp_name:
            database_path = Path(temp_name) / "registry.sqlite"
            sqlite3.connect(database_path).close()
            MODULE.BASE.ensure_registry_schema(database_path)
            result = MODULE.ensure_flexible_model_registry_schema(database_path)
            self.assertEqual(result["status"], "migrated")
            with sqlite3.connect(database_path) as connection:
                event_sql = connection.execute(
                    "SELECT sql FROM sqlite_master WHERE name='model_events'"
                ).fetchone()[0]
                artifact_sql = connection.execute(
                    "SELECT sql FROM sqlite_master WHERE name='model_artifacts'"
                ).fetchone()[0]
                self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone()[0], "ok")
            self.assertIn("output_artifacts_count > 0", event_sql)
            self.assertIn("length(trim(artifact_role)) > 0", artifact_sql)

    def test_pipeline_contains_rfecv_and_final_forest(self) -> None:
        matrix, labels, _ = synthetic_data()
        pipeline = MODULE.build_pipeline(
            None,
            minimum_features=2,
            rfecv_step=0.25,
            inner_cv_splits=3,
            n_estimators=10,
            estimator_n_jobs=1,
            selector_n_jobs=1,
        )
        self.assertEqual(list(pipeline.named_steps), ["feature_selection", "classifier"])
        pipeline.fit(matrix, labels)
        selector = pipeline.named_steps["feature_selection"]
        self.assertGreaterEqual(selector.n_features_, 2)
        self.assertEqual(selector.n_features_in_, 12)
        self.assertEqual(
            pipeline.named_steps["classifier"].n_features_in_, selector.n_features_
        )

    def test_nested_cv_refits_selector_inside_each_outer_fold(self) -> None:
        matrix, labels, _ = synthetic_data()
        outer = list(
            StratifiedKFold(n_splits=3, shuffle=True, random_state=42).split(
                np.zeros(len(labels)), labels
            )
        )
        results, summaries, curves = MODULE.run_nested_cross_validation(
            matrix,
            labels,
            outer,
            minimum_features=2,
            rfecv_step=0.25,
            inner_cv_splits=3,
            n_estimators=10,
            estimator_n_jobs=1,
            selector_n_jobs=1,
        )
        self.assertEqual(set(summaries), {"C1", "C2"})
        self.assertEqual(len(results.loc[results["record_type"].eq("fold")]), 6)
        self.assertEqual(
            curves.groupby(["configuration_id", "outer_fold"]).ngroups, 6
        )
        self.assertTrue(
            curves.groupby(["configuration_id", "outer_fold"])[
                "number_of_retained_features"
            ]
            .min()
            .eq(2)
            .all()
        )
        self.assertTrue(
            curves.groupby(["configuration_id", "outer_fold"])[
                "selected_as_optimal"
            ]
            .sum()
            .eq(1)
            .all()
        )

    def test_tie_rule_selects_c1(self) -> None:
        summaries = {
            "C1": {"mean_macro_f1": 0.7000000000001},
            "C2": {"mean_macro_f1": 0.7000000000002},
        }
        self.assertEqual(MODULE.select_configuration(summaries, 12), "C1")
        summaries["C2"]["mean_macro_f1"] = 0.700001
        self.assertEqual(MODULE.select_configuration(summaries, 12), "C2")

    def test_cross_family_missing_values_become_valid_json_nulls(self) -> None:
        frame = pd.DataFrame(
            [
                {"configuration_id": "A1", "rfecv_selected_features": np.nan},
                {"configuration_id": "C1", "rfecv_selected_features": 124.0},
            ]
        )
        records = MODULE.dataframe_records_for_json(frame)
        text = MODULE.BASE.json_text(records)
        self.assertIsNone(json.loads(text)[0]["rfecv_selected_features"])
        with sqlite3.connect(":memory:") as connection:
            self.assertEqual(connection.execute("SELECT json_valid(?)", (text,)).fetchone()[0], 1)

    def test_feature_curve_plot_and_workbook_embed_training_results(self) -> None:
        matrix, labels, features = synthetic_data()
        curves = []
        pipelines = {}
        for configuration_id, configuration in MODULE.CONFIGURATIONS.items():
            pipeline = MODULE.build_pipeline(
                configuration["class_weight"],
                minimum_features=2,
                rfecv_step=0.25,
                inner_cv_splits=3,
                n_estimators=10,
                estimator_n_jobs=1,
                selector_n_jobs=1,
            )
            pipeline.fit(matrix, labels)
            pipelines[configuration_id] = pipeline
            curves.append(
                MODULE.extract_rfecv_curve(
                    pipeline.named_steps["feature_selection"],
                    configuration_id=configuration_id,
                    fit_scope="complete_training",
                    outer_fold=None,
                    minimum_features=2,
                    rfecv_step=0.25,
                )
            )
        curve_data = pd.concat(curves, ignore_index=True)
        selected_pipeline = pipelines["C1"]
        selected_features = MODULE.build_selected_features(features, selected_pipeline)
        metric_tables = []
        class_tables = []
        confusion_tables = []
        prediction_tables = []
        for configuration_id in MODULE.CONFIGURATIONS:
            _, metrics, classes, matrix_table, configuration_predictions = (
                MODULE.build_evaluation_outputs(
                    configuration_id,
                    pipelines[configuration_id],
                    matrix[:6],
                    synthetic_index(labels[:6]),
                )
            )
            metric_tables.append(metrics)
            class_tables.append(classes)
            confusion_tables.append(matrix_table)
            prediction_tables.append(configuration_predictions)
        metrics = pd.concat(metric_tables, ignore_index=True)
        per_class = pd.concat(class_tables, ignore_index=True)
        confusion = pd.concat(confusion_tables, ignore_index=True)
        predictions = pd.concat(prediction_tables, ignore_index=True)
        self.assertEqual(set(predictions["configuration_id"]), {"C1", "C2"})
        self.assertEqual(len(predictions), 12)
        workbook_sheets = {
            "Model_Summary": pd.DataFrame(
                [("identity", "model_id", "MODEL_C_TEST", "Synthetic")],
                columns=["section", "item", "value", "description"],
            ),
            "CV_Comparison": pd.DataFrame(
                [
                    {
                        "configuration_id": configuration_id,
                        "record_type": record_type,
                        "fold": fold,
                        "selected_configuration": configuration_id == "C1",
                    }
                    for configuration_id in ("A1", "A2", "B1", "B2", "C1", "C2")
                    for record_type, fold in [("fold", 1), ("mean", "all"), ("std", "all")]
                ]
            ),
            "Test_Metrics": metrics,
            "Per_Class_Metrics": per_class,
            "Confusion_Matrix": confusion,
            "Selected_Features": selected_features,
            "Test_Predictions": predictions,
            "RFECV_Feature_Curve": curve_data,
            "Artifact_Descriptions": MODULE.artifact_descriptions_table(),
        }
        with tempfile.TemporaryDirectory() as temp_name:
            output_dir = Path(temp_name)
            plot_path = output_dir / MODULE.FEATURE_CURVE_PLOT_FILENAME
            MODULE.plot_feature_curve(
                curve_data,
                plot_path,
                minimum_features=2,
                model_a_features=4,
            )
            with Image.open(plot_path) as image:
                self.assertEqual(image.format, "PNG")
                self.assertGreaterEqual(image.width, 1200)
                self.assertGreaterEqual(image.height, 700)
            workbook_path = output_dir / MODULE.PERFORMANCE_REPORT_FILENAME
            MODULE.write_performance_workbook(
                workbook_path, workbook_sheets, plot_path
            )
            workbook = load_workbook(workbook_path, read_only=False)
            try:
                self.assertEqual(workbook.sheetnames, list(workbook_sheets))
                self.assertEqual(len(workbook["RFECV_Feature_Curve"]._images), 1)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
