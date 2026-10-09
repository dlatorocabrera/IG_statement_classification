#!/usr/bin/env python3
"""Train and register proportional Wroblewska-inspired Random Forest Model A."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sqlite3
import sys
import tempfile
import warnings
from collections import Counter
from datetime import datetime, timezone
from functools import partial
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd
import sklearn
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from scipy import sparse
from sklearn.ensemble import RandomForestClassifier
from sklearn.feature_selection import SelectKBest, mutual_info_classif
from sklearn.metrics import (
    accuracy_score,
    confusion_matrix,
    f1_score,
    precision_recall_fscore_support,
    precision_score,
    recall_score,
)
from sklearn.model_selection import StratifiedKFold
from sklearn.pipeline import Pipeline


SCRIPT_DIR = Path(__file__).resolve().parent
CODES_DIR = SCRIPT_DIR.parent
PROJECT_ROOT = CODES_DIR.parent

if str(CODES_DIR) not in sys.path:
    sys.path.insert(0, str(CODES_DIR))

try:
    from config import DATABASE_PATH
except ImportError as error:
    raise SystemExit("Could not import codes/config.py.") from error

DEFAULT_OUTPUT_DIR = (
    PROJECT_ROOT / "outputs" / "task_3_statement_classification" / "models"
)

SCRIPT_NAME = "09_train_random_forest_model_a.py"
SCRIPT_VERSION = "1.1"
MODEL_NAME = "Model A: paired proportional Wroblewska-inspired Random Forests"
MODEL_DESIGNATION = "Model A"
MODEL_DESCRIPTION = (
    "Paired three-class Random Forest evaluation comparing unweighted A1 and "
    "class-balanced A2. Both use mutual-information selection retaining the same "
    "vocabulary proportion as the 70-of-3,770 Wroblewska reference pipeline."
)

RANDOM_STATE = 42
CLASS_ORDER = ("regulative", "constitutive", "non_institutional")
REFERENCE_VOCABULARY_SIZE = 3_770
REFERENCE_SELECTED_FEATURE_COUNT = 70
REFERENCE_RETENTION_PROPORTION = (
    REFERENCE_SELECTED_FEATURE_COUNT / REFERENCE_VOCABULARY_SIZE
)
EXPECTED_CURRENT_VOCABULARY_SIZE = 23_464
EXPECTED_SELECTED_FEATURE_COUNT = round(
    EXPECTED_CURRENT_VOCABULARY_SIZE * REFERENCE_RETENTION_PROPORTION
)
CV_SPLITS = 5
STORED_METRIC_DECIMALS = 12

MODEL_FILENAME = "model_a_proportional_rf.joblib"
CV_RESULTS_FILENAME = "model_a_cv_results.csv"
SELECTED_FEATURES_FILENAME = "model_a_selected_features.csv"
TEST_PREDICTIONS_FILENAME = "model_a_test_predictions.csv"
CONFUSION_MATRIX_FILENAME = "model_a_confusion_matrix.csv"
PERFORMANCE_REPORT_FILENAME = "model_a_proportional_rf_performance.xlsx"

REQUIRED_VECTOR_ARTIFACTS = {
    "training_matrix": "X_train_tfidf.npz",
    "test_matrix": "X_test_tfidf.npz",
    "training_index": "train_vectorization_index.csv",
    "test_index": "test_vectorization_index.csv",
    "features": "tfidf_features.csv",
    "fitted_vectorizer": "tfidf_vectorizer.joblib",
}
INDEX_COLUMNS = (
    "matrix_row_index",
    "candidate_id",
    "candidate_text_original",
    "reviewer_gold_label",
    "split",
)
FEATURE_COLUMNS = ("feature_index", "ngram_text", "ngram_length", "idf_value")

CONFIGURATIONS = {
    "A1": {"class_weight": None, "label": "unweighted baseline"},
    "A2": {"class_weight": "balanced", "label": "balanced adaptation"},
}
RANDOM_FOREST_PARAMETERS = {
    "n_estimators": 100,
    "criterion": "gini",
    "max_depth": None,
    "min_samples_split": 2,
    "min_samples_leaf": 1,
    "max_features": "sqrt",
    "bootstrap": True,
    "random_state": RANDOM_STATE,
    "n_jobs": -1,
}

ARTIFACT_DESCRIPTIONS = {
    "fitted_model": (
        "Fitted-model bundle containing the final 436-feature A1 and A2 pipelines, "
        "configuration metadata, class order, and CV-selected primary reference."
    ),
    "cross_validation_results": (
        "Fold-level, mean, and standard-deviation metrics for configurations A1 and A2."
    ),
    "selected_features": (
        "Final 436 training-selected TF-IDF features, ranks, source indices, and "
        "mutual-information scores."
    ),
    "test_predictions": (
        "Long-form candidate-aligned held-out predictions and correctness flags for "
        "both A1 and A2."
    ),
    "confusion_matrix": (
        "Held-out confusion matrices for A1 and A2 using the predefined class order."
    ),
    "performance_report": (
        "Human-readable workbook containing lineage, CV comparison, test performance, "
        "selected features, predictions, and artifact descriptions."
    ),
}


MODEL_REGISTRY_SCHEMA = """
CREATE TABLE IF NOT EXISTS model_events (
    model_id TEXT PRIMARY KEY,
    vectorization_id TEXT NOT NULL,
    split_id TEXT NOT NULL,
    matching_id TEXT NOT NULL,
    segmentation_id TEXT NOT NULL,
    document_id TEXT NOT NULL,
    model_date TEXT NOT NULL,
    created_at TEXT NOT NULL,
    model_name TEXT NOT NULL,
    model_designation TEXT NOT NULL,
    model_description TEXT NOT NULL,
    script_path TEXT NOT NULL,
    script_version TEXT NOT NULL,
    model_tool TEXT NOT NULL,
    model_tool_version TEXT NOT NULL,
    input_artifacts_json TEXT NOT NULL,
    input_artifacts_count INTEGER NOT NULL,
    class_order_json TEXT NOT NULL,
    class_distributions_json TEXT NOT NULL,
    training_records_count INTEGER NOT NULL,
    test_records_count INTEGER NOT NULL,
    reference_vocabulary_size INTEGER NOT NULL,
    reference_selected_feature_count INTEGER NOT NULL,
    retention_proportion REAL NOT NULL,
    retention_calculation TEXT NOT NULL,
    current_vocabulary_size INTEGER NOT NULL,
    current_selected_feature_count INTEGER NOT NULL,
    feature_selection_method TEXT NOT NULL,
    feature_selection_parameters_json TEXT NOT NULL,
    cross_validation_design_json TEXT NOT NULL,
    random_forest_configurations_json TEXT NOT NULL,
    cross_validation_results_json TEXT NOT NULL,
    configuration_selection_rule TEXT NOT NULL,
    selected_configuration TEXT NOT NULL,
    final_test_metrics_json TEXT NOT NULL,
    validation_results_json TEXT NOT NULL,
    output_artifacts_count INTEGER NOT NULL,
    model_status TEXT NOT NULL,
    notes TEXT,
    CHECK (length(model_date) = 10),
    CHECK (input_artifacts_count = 6),
    CHECK (training_records_count > 0),
    CHECK (test_records_count > 0),
    CHECK (reference_vocabulary_size > 0),
    CHECK (reference_selected_feature_count > 0),
    CHECK (retention_proportion > 0 AND retention_proportion <= 1),
    CHECK (current_vocabulary_size >= current_selected_feature_count),
    CHECK (current_selected_feature_count > 0),
    CHECK (selected_configuration IN ('A1', 'A2')),
    CHECK (output_artifacts_count = 6),
    CHECK (json_valid(input_artifacts_json)),
    CHECK (json_valid(class_order_json)),
    CHECK (json_valid(class_distributions_json)),
    CHECK (json_valid(feature_selection_parameters_json)),
    CHECK (json_valid(cross_validation_design_json)),
    CHECK (json_valid(random_forest_configurations_json)),
    CHECK (json_valid(cross_validation_results_json)),
    CHECK (json_valid(final_test_metrics_json)),
    CHECK (json_valid(validation_results_json)),
    FOREIGN KEY (vectorization_id)
        REFERENCES vectorization_events(vectorization_id)
        ON UPDATE CASCADE ON DELETE RESTRICT,
    FOREIGN KEY (split_id)
        REFERENCES split_events(split_id)
        ON UPDATE CASCADE ON DELETE RESTRICT,
    FOREIGN KEY (matching_id)
        REFERENCES matching_events(matching_id)
        ON UPDATE CASCADE ON DELETE RESTRICT,
    FOREIGN KEY (segmentation_id)
        REFERENCES segmentation_events(segmentation_id)
        ON UPDATE CASCADE ON DELETE RESTRICT,
    FOREIGN KEY (document_id)
        REFERENCES document_inventory(document_id)
        ON UPDATE CASCADE ON DELETE RESTRICT
);

CREATE TABLE IF NOT EXISTS model_artifacts (
    model_id TEXT NOT NULL,
    artifact_role TEXT NOT NULL,
    file_name TEXT NOT NULL,
    artifact_path TEXT NOT NULL,
    artifact_hash_sha256 TEXT NOT NULL,
    size_bytes INTEGER NOT NULL,
    artifact_description TEXT NOT NULL,
    created_at TEXT NOT NULL,
    PRIMARY KEY (model_id, artifact_role),
    UNIQUE (model_id, file_name),
    CHECK (artifact_role IN (
        'fitted_model', 'cross_validation_results', 'selected_features',
        'test_predictions', 'confusion_matrix', 'performance_report'
    )),
    CHECK (size_bytes > 0),
    FOREIGN KEY (model_id)
        REFERENCES model_events(model_id)
        ON UPDATE CASCADE ON DELETE RESTRICT
);

CREATE INDEX IF NOT EXISTS idx_model_events_vectorization_created
    ON model_events(vectorization_id, created_at);

CREATE INDEX IF NOT EXISTS idx_model_artifacts_path
    ON model_artifacts(artifact_path);
"""


def utc_now_iso() -> str:
    """Return a UTC timestamp without microseconds."""
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def project_relative(path: Path) -> str:
    """Return a project-relative path when possible."""
    try:
        return path.resolve().relative_to(PROJECT_ROOT.resolve()).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def resolve_project_path(value: str | Path) -> Path:
    """Resolve an absolute or project-relative path."""
    path = Path(value).expanduser()
    return path.resolve() if path.is_absolute() else (PROJECT_ROOT / path).resolve()


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    """Calculate a file SHA-256 hash."""
    digest = hashlib.sha256()
    with path.open("rb") as file:
        while chunk := file.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def json_default(value: Any) -> Any:
    """Convert numpy scalar values or reject unsupported JSON objects."""
    if isinstance(value, np.generic):
        return value.item()
    raise TypeError(f"Object of type {type(value).__name__} is not JSON serializable")


def json_text(value: Any) -> str:
    """Serialize JSON deterministically for SQLite."""
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=json_default,
    )


def ensure_registry_schema(database_path: Path) -> None:
    """Create append-only model registry tables when absent."""
    if not database_path.exists():
        raise FileNotFoundError(f"Corpus inventory database not found: {database_path}")
    with sqlite3.connect(database_path) as connection:
        connection.execute("PRAGMA foreign_keys = ON")
        connection.executescript(MODEL_REGISTRY_SCHEMA)


def next_model_id(database_path: Path) -> str:
    """Return the next append-only MODEL_###### identifier."""
    with sqlite3.connect(database_path) as connection:
        identifiers = [
            str(row[0])
            for row in connection.execute("SELECT model_id FROM model_events")
        ]
    invalid = [value for value in identifiers if not re.fullmatch(r"MODEL_\d{6}", value)]
    if invalid:
        raise ValueError("Invalid model identifiers: " + ", ".join(invalid[:10]))
    highest = max([0] + [int(value.removeprefix("MODEL_")) for value in identifiers])
    return f"MODEL_{highest + 1:06d}"


def load_vectorization_event(
    database_path: Path,
    vectorization_id: str | None,
) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    """Load one completed vectorization event, its lineage, and six artifacts."""
    with sqlite3.connect(database_path) as connection:
        connection.row_factory = sqlite3.Row
        if vectorization_id is None:
            row = connection.execute(
                """
                SELECT * FROM vectorization_events
                WHERE vectorization_status IN ('Completed', 'Completed with warnings')
                ORDER BY created_at DESC, vectorization_id DESC
                LIMIT 1
                """
            ).fetchone()
        else:
            row = connection.execute(
                "SELECT * FROM vectorization_events WHERE vectorization_id = ?",
                (vectorization_id,),
            ).fetchone()
        if row is None:
            requested = vectorization_id or "latest completed vectorization"
            raise ValueError(f"Registered vectorization event not found: {requested}")
        event = dict(row)
        if event["vectorization_status"] not in {
            "Completed",
            "Completed with warnings",
        }:
            raise ValueError(f"Vectorization {event['vectorization_id']} is not completed.")

        split = connection.execute(
            "SELECT * FROM split_events WHERE split_id = ?", (event["split_id"],)
        ).fetchone()
        if split is None:
            raise ValueError("The vectorization event has no registered split lineage.")
        matching = connection.execute(
            "SELECT * FROM matching_events WHERE matching_id = ?",
            (event["matching_id"],),
        ).fetchone()
        if matching is None:
            raise ValueError("The vectorization event has no registered matching lineage.")
        for ancestor, name in ((split, "split"), (matching, "matching")):
            for field in ("matching_id", "segmentation_id", "document_id"):
                if str(event[field]) != str(ancestor[field]):
                    raise ValueError(
                        f"Vectorization/{name} lineage mismatch for {field}: "
                        f"{event[field]!r} != {ancestor[field]!r}."
                    )

        artifact_rows = connection.execute(
            """
            SELECT * FROM vectorization_artifacts
            WHERE vectorization_id = ?
            ORDER BY artifact_role
            """,
            (event["vectorization_id"],),
        ).fetchall()
    artifacts = {str(item["artifact_role"]): dict(item) for item in artifact_rows}
    expected_roles = set(REQUIRED_VECTOR_ARTIFACTS)
    if set(artifacts) != expected_roles:
        raise ValueError(
            "Vectorization artifact roles differ from the required set: "
            f"found={sorted(artifacts)}, expected={sorted(expected_roles)}."
        )
    for role, expected_name in REQUIRED_VECTOR_ARTIFACTS.items():
        if artifacts[role]["file_name"] != expected_name:
            raise ValueError(
                f"Unexpected {role} filename: {artifacts[role]['file_name']!r}."
            )
    return event, artifacts


def class_counts(labels: pd.Series | np.ndarray) -> dict[str, int]:
    """Count labels in the explicit protocol order."""
    counts = Counter(str(value) for value in labels)
    return {label: int(counts.get(label, 0)) for label in CLASS_ORDER}


def load_and_validate_inputs(
    event: dict[str, Any],
    artifacts: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    """Read and validate registered matrices, indices, features, and vectorizer."""
    paths = {
        role: resolve_project_path(values["artifact_path"])
        for role, values in artifacts.items()
    }
    checks: dict[str, bool] = {}
    for role, path in paths.items():
        checks[f"{role}_exists"] = path.is_file()
        if not path.is_file():
            raise FileNotFoundError(f"Registered vectorization artifact not found: {path}")
        checks[f"{role}_hash_matches_registry"] = (
            sha256_file(path) == artifacts[role]["artifact_hash_sha256"]
        )
        checks[f"{role}_size_matches_registry"] = (
            path.stat().st_size == int(artifacts[role]["size_bytes"])
        )

    train_matrix = sparse.load_npz(paths["training_matrix"]).tocsr()
    test_matrix = sparse.load_npz(paths["test_matrix"]).tocsr()
    train_index = pd.read_csv(
        paths["training_index"], dtype=object, keep_default_na=False, encoding="utf-8"
    )
    test_index = pd.read_csv(
        paths["test_index"], dtype=object, keep_default_na=False, encoding="utf-8"
    )
    features = pd.read_csv(
        paths["features"], dtype=object, keep_default_na=False, encoding="utf-8"
    )
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Trying to unpickle estimator .* from version .*",
            category=UserWarning,
        )
        vectorizer = joblib.load(paths["fitted_vectorizer"])

    checks.update(
        {
            "training_matrix_shape_matches_registry": train_matrix.shape
            == (
                int(event["training_matrix_rows"]),
                int(event["training_matrix_columns"]),
            ),
            "test_matrix_shape_matches_registry": test_matrix.shape
            == (
                int(event["test_matrix_rows"]),
                int(event["test_matrix_columns"]),
            ),
            "training_matrix_nonzero_matches_registry": train_matrix.nnz
            == int(event["training_matrix_nonzero"]),
            "test_matrix_nonzero_matches_registry": test_matrix.nnz
            == int(event["test_matrix_nonzero"]),
            "training_matrix_values_finite": bool(np.isfinite(train_matrix.data).all()),
            "test_matrix_values_finite": bool(np.isfinite(test_matrix.data).all()),
            "training_index_columns_valid": list(train_index.columns)
            == list(INDEX_COLUMNS),
            "test_index_columns_valid": list(test_index.columns) == list(INDEX_COLUMNS),
            "feature_columns_valid": list(features.columns) == list(FEATURE_COLUMNS),
            "training_index_rows_match_matrix": len(train_index)
            == train_matrix.shape[0],
            "test_index_rows_match_matrix": len(test_index) == test_matrix.shape[0],
            "feature_rows_match_matrix": len(features) == train_matrix.shape[1],
            "vocabulary_at_least_selected_count": train_matrix.shape[1]
            >= EXPECTED_SELECTED_FEATURE_COUNT,
        }
    )

    def validate_index(index: pd.DataFrame, partition: str) -> dict[str, bool]:
        candidate_ids = index["candidate_id"].map(str)
        labels = index["reviewer_gold_label"].map(str)
        texts = index["candidate_text_original"]
        return {
            "row_indices_sequential": [int(value) for value in index["matrix_row_index"]]
            == list(range(len(index))),
            "candidate_ids_nonblank": bool(candidate_ids.str.strip().ne("").all()),
            "candidate_ids_unique": not candidate_ids.duplicated().any(),
            "texts_are_strings": bool(texts.map(lambda value: isinstance(value, str)).all()),
            "texts_nonblank": bool(texts.map(str).str.strip().ne("").all()),
            "labels_approved": bool(labels.isin(CLASS_ORDER).all()),
            "partition_preserved": bool(index["split"].map(str).eq(partition).all()),
        }

    for prefix, result in (
        ("training", validate_index(train_index, "training")),
        ("test", validate_index(test_index, "test")),
    ):
        checks.update({f"{prefix}_{name}": value for name, value in result.items()})

    checks["candidate_ids_do_not_cross_partitions"] = not (
        set(train_index["candidate_id"].map(str))
        & set(test_index["candidate_id"].map(str))
    )
    checks["feature_indices_sequential"] = [
        int(value) for value in features["feature_index"]
    ] == list(range(len(features)))
    checks["ngram_lengths_valid"] = [
        int(value) for value in features["ngram_length"]
    ] == [str(value).count(" ") + 1 for value in features["ngram_text"]]
    checks["feature_idf_values_finite"] = bool(
        np.isfinite(features["idf_value"].astype(float).to_numpy()).all()
    )
    vectorizer_features = vectorizer.get_feature_names_out().tolist()
    checks["feature_order_matches_fitted_vectorizer"] = (
        features["ngram_text"].map(str).tolist() == vectorizer_features
    )
    checks["feature_idf_matches_fitted_vectorizer"] = bool(
        np.allclose(
            features["idf_value"].astype(float).to_numpy(),
            np.asarray(vectorizer.idf_, dtype=float),
            rtol=0,
            atol=1e-12,
        )
    )

    observed_distributions = {
        "training": class_counts(train_index["reviewer_gold_label"]),
        "test": class_counts(test_index["reviewer_gold_label"]),
    }
    registered_distributions = json.loads(event["class_distributions_json"])
    checks["class_distributions_match_vectorization_event"] = all(
        observed_distributions[partition]
        == {
            label: int(registered_distributions[partition][label])
            for label in CLASS_ORDER
        }
        for partition in ("training", "test")
    )
    checks["training_row_count_matches_vectorization_event"] = len(train_index) == int(
        event["input_training_rows"]
    )
    checks["test_row_count_matches_vectorization_event"] = len(test_index) == int(
        event["input_test_rows"]
    )
    checks["vocabulary_matches_vectorization_event"] = train_matrix.shape[1] == int(
        event["vocabulary_size"]
    )

    checks = {name: bool(value) for name, value in checks.items()}
    failed = [name for name, passed in checks.items() if not passed]
    if failed:
        raise ValueError("Model input validation failed: " + ", ".join(failed))
    return {
        "train_matrix": train_matrix,
        "test_matrix": test_matrix,
        "train_index": train_index,
        "test_index": test_index,
        "features": features,
        "paths": paths,
        "class_distributions": observed_distributions,
        "validation": {"status": "passed", "checks": checks},
    }


def selected_feature_count(vocabulary_size: int) -> int:
    """Scale the reference 70-of-3,770 retention rate to a vocabulary size."""
    return round(vocabulary_size * REFERENCE_RETENTION_PROPORTION)


def reproducible_mutual_information(
    matrix: sparse.spmatrix | np.ndarray,
    labels: np.ndarray,
) -> np.ndarray:
    """Calculate deterministic MI scores while silencing one known sklearn flood."""
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message=(
                "Clustering metrics expects discrete values but received continuous "
                "values for label, and multiclass values for target"
            ),
            category=UserWarning,
        )
        return mutual_info_classif(matrix, labels, random_state=RANDOM_STATE)


def build_pipeline(class_weight: str | None, k_features: int) -> Pipeline:
    """Build the leakage-safe selector/classifier pipeline."""
    return Pipeline(
        [
            (
                "feature_selection",
                SelectKBest(
                    score_func=reproducible_mutual_information,
                    k=k_features,
                ),
            ),
            (
                "classifier",
                RandomForestClassifier(
                    **RANDOM_FOREST_PARAMETERS,
                    class_weight=class_weight,
                ),
            ),
        ]
    )


def build_cv_folds(
    labels: np.ndarray,
) -> tuple[list[tuple[np.ndarray, np.ndarray]], dict[str, Any]]:
    """Materialize and validate the identical folds used by both configurations."""
    splitter = StratifiedKFold(
        n_splits=CV_SPLITS,
        shuffle=True,
        random_state=RANDOM_STATE,
    )
    folds = list(splitter.split(np.zeros(len(labels), dtype=np.int8), labels))
    checks: dict[str, bool] = {"five_folds_materialized": len(folds) == CV_SPLITS}
    fold_distributions: list[dict[str, Any]] = []
    required = set(CLASS_ORDER)
    for fold_number, (training_rows, validation_rows) in enumerate(folds, start=1):
        train_set = set(labels[training_rows])
        validation_set = set(labels[validation_rows])
        checks[f"fold_{fold_number}_training_contains_all_classes"] = train_set == required
        checks[f"fold_{fold_number}_validation_contains_all_classes"] = (
            validation_set == required
        )
        checks[f"fold_{fold_number}_partitions_do_not_overlap"] = not (
            set(training_rows.tolist()) & set(validation_rows.tolist())
        )
        checks[f"fold_{fold_number}_covers_all_training_rows"] = (
            len(training_rows) + len(validation_rows) == len(labels)
        )
        fold_distributions.append(
            {
                "fold": fold_number,
                "training": class_counts(labels[training_rows]),
                "validation": class_counts(labels[validation_rows]),
            }
        )
    checks = {name: bool(value) for name, value in checks.items()}
    failed = [name for name, passed in checks.items() if not passed]
    if failed:
        raise ValueError("Cross-validation fold validation failed: " + ", ".join(failed))
    return folds, {
        "status": "passed",
        "checks": checks,
        "fold_class_distributions": fold_distributions,
    }


def calculate_metrics(actual: np.ndarray, predicted: np.ndarray) -> dict[str, float]:
    """Calculate the five protocol metrics using the fixed class order."""
    return {
        "macro_f1": float(
            f1_score(actual, predicted, labels=CLASS_ORDER, average="macro", zero_division=0)
        ),
        "weighted_f1": float(
            f1_score(
                actual,
                predicted,
                labels=CLASS_ORDER,
                average="weighted",
                zero_division=0,
            )
        ),
        "accuracy": float(accuracy_score(actual, predicted)),
        "macro_precision": float(
            precision_score(
                actual,
                predicted,
                labels=CLASS_ORDER,
                average="macro",
                zero_division=0,
            )
        ),
        "macro_recall": float(
            recall_score(
                actual,
                predicted,
                labels=CLASS_ORDER,
                average="macro",
                zero_division=0,
            )
        ),
    }


def run_cross_validation(
    matrix: sparse.csr_matrix,
    labels: np.ndarray,
    folds: list[tuple[np.ndarray, np.ndarray]],
    k_features: int,
) -> tuple[pd.DataFrame, dict[str, dict[str, Any]]]:
    """Evaluate A1 and A2 on the same materialized folds."""
    metric_names = (
        "macro_f1",
        "weighted_f1",
        "accuracy",
        "macro_precision",
        "macro_recall",
    )
    fold_rows: list[dict[str, Any]] = []
    summaries: dict[str, dict[str, Any]] = {}
    for configuration_id, configuration in CONFIGURATIONS.items():
        configuration_rows: list[dict[str, Any]] = []
        for fold_number, (training_rows, validation_rows) in enumerate(folds, start=1):
            pipeline = build_pipeline(configuration["class_weight"], k_features)
            pipeline.fit(matrix[training_rows], labels[training_rows])
            selector = pipeline.named_steps["feature_selection"]
            if int(selector.get_support().sum()) != k_features:
                raise RuntimeError(
                    f"{configuration_id} fold {fold_number} did not select {k_features} features."
                )
            predicted = pipeline.predict(matrix[validation_rows])
            metrics = calculate_metrics(labels[validation_rows], predicted)
            row = {
                "configuration_id": configuration_id,
                "configuration_label": configuration["label"],
                "class_weight": configuration["class_weight"] or "None",
                "record_type": "fold",
                "fold": fold_number,
                "training_records": len(training_rows),
                "validation_records": len(validation_rows),
                **metrics,
                "selected_configuration": False,
            }
            fold_rows.append(row)
            configuration_rows.append(row)

        summary: dict[str, Any] = {}
        for metric in metric_names:
            values = np.asarray([row[metric] for row in configuration_rows], dtype=float)
            summary[f"mean_{metric}"] = float(values.mean())
            summary[f"std_{metric}"] = float(values.std(ddof=1))
        summaries[configuration_id] = summary
        for record_type, prefix in (("mean", "mean"), ("std", "std")):
            fold_rows.append(
                {
                    "configuration_id": configuration_id,
                    "configuration_label": configuration["label"],
                    "class_weight": configuration["class_weight"] or "None",
                    "record_type": record_type,
                    "fold": "all",
                    "training_records": "",
                    "validation_records": "",
                    **{
                        metric: summary[f"{prefix}_{metric}"] for metric in metric_names
                    },
                    "selected_configuration": False,
                }
            )
    return pd.DataFrame(fold_rows), summaries


def select_configuration(
    summaries: dict[str, dict[str, Any]],
    decimal_places: int = STORED_METRIC_DECIMALS,
) -> str:
    """Choose the highest mean macro-F1, selecting A1 on a stored-precision tie."""
    a1 = round(float(summaries["A1"]["mean_macro_f1"]), decimal_places)
    a2 = round(float(summaries["A2"]["mean_macro_f1"]), decimal_places)
    return "A1" if a1 >= a2 else "A2"


def build_selected_features(
    pipeline: Pipeline,
    features: pd.DataFrame,
) -> pd.DataFrame:
    """Map the final selector support and MI scores to the TF-IDF feature table."""
    selector: SelectKBest = pipeline.named_steps["feature_selection"]
    selected_indices = selector.get_support(indices=True).astype(int)
    scores = np.asarray(selector.scores_, dtype=float)[selected_indices]
    if not np.isfinite(scores).all():
        raise RuntimeError("Selected mutual-information scores contain non-finite values.")
    selected = pd.DataFrame(
        {
            "original_tfidf_feature_index": selected_indices,
            "mutual_information_score": scores,
        }
    )
    feature_lookup = features.copy()
    feature_lookup["feature_index"] = feature_lookup["feature_index"].astype(int)
    feature_lookup["ngram_length"] = feature_lookup["ngram_length"].astype(int)
    feature_lookup["idf_value"] = feature_lookup["idf_value"].astype(float)
    selected = selected.merge(
        feature_lookup,
        how="left",
        left_on="original_tfidf_feature_index",
        right_on="feature_index",
        validate="one_to_one",
    ).drop(columns="feature_index")
    selected = selected.sort_values(
        ["mutual_information_score", "original_tfidf_feature_index"],
        ascending=[False, True],
        kind="mergesort",
    ).reset_index(drop=True)
    selected.insert(0, "final_relevance_rank", np.arange(1, len(selected) + 1))
    return selected.loc[
        :,
        [
            "final_relevance_rank",
            "original_tfidf_feature_index",
            "ngram_text",
            "ngram_length",
            "mutual_information_score",
            "idf_value",
        ],
    ]


def build_evaluation_outputs(
    configuration_id: str,
    pipeline: Pipeline,
    matrix: sparse.csr_matrix,
    test_index: pd.DataFrame,
) -> tuple[dict[str, Any], pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Evaluate one predefined configuration and construct report-ready tables."""
    if configuration_id not in CONFIGURATIONS:
        raise ValueError(f"Unknown Model A configuration: {configuration_id}")
    configuration_label = CONFIGURATIONS[configuration_id]["label"]
    actual = test_index["reviewer_gold_label"].map(str).to_numpy()
    predicted = pipeline.predict(matrix).astype(str)
    overall = calculate_metrics(actual, predicted)
    correct = actual == predicted
    overall.update(
        {
            "correct_count": int(correct.sum()),
            "incorrect_count": int((~correct).sum()),
            "correct_proportion": float(correct.mean()),
            "incorrect_proportion": float((~correct).mean()),
            "test_records": int(len(actual)),
        }
    )

    precision, recall, f1, support = precision_recall_fscore_support(
        actual,
        predicted,
        labels=CLASS_ORDER,
        zero_division=0,
    )
    per_class = pd.DataFrame(
        {
            "configuration_id": configuration_id,
            "configuration_label": configuration_label,
            "class": CLASS_ORDER,
            "precision": precision,
            "recall": recall,
            "f1": f1,
            "support": support.astype(int),
        }
    )
    matrix_values = confusion_matrix(actual, predicted, labels=CLASS_ORDER)
    confusion = pd.DataFrame(
        matrix_values,
        columns=[f"predicted_{label}" for label in CLASS_ORDER],
    )
    confusion.insert(0, "configuration_label", configuration_label)
    confusion.insert(0, "configuration_id", configuration_id)
    confusion.insert(2, "actual_label", CLASS_ORDER)
    confusion["actual_total"] = matrix_values.sum(axis=1)

    predictions = pd.DataFrame(
        {
            "configuration_id": configuration_id,
            "configuration_label": configuration_label,
            "matrix_row_index": test_index["matrix_row_index"].astype(int),
            "candidate_id": test_index["candidate_id"].map(str),
            "candidate_text_original": test_index["candidate_text_original"].map(str),
            "gold_label": actual,
            "predicted_label": predicted,
            "correct": correct,
        }
    )
    metric_rows = [
            {
                "metric": "macro_f1",
                "value": overall["macro_f1"],
                "role": "principal metric",
                "description": "Unweighted mean of the three class-specific F1 scores.",
            },
            {
                "metric": "weighted_f1",
                "value": overall["weighted_f1"],
                "role": "secondary metric",
                "description": "Class-support-weighted mean of class-specific F1 scores.",
            },
            {
                "metric": "accuracy",
                "value": overall["accuracy"],
                "role": "secondary metric",
                "description": "Proportion of held-out statements classified correctly.",
            },
            {
                "metric": "macro_precision",
                "value": overall["macro_precision"],
                "role": "secondary metric",
                "description": "Unweighted mean precision across the three classes.",
            },
            {
                "metric": "macro_recall",
                "value": overall["macro_recall"],
                "role": "secondary metric",
                "description": "Unweighted mean recall across the three classes.",
            },
            {
                "metric": "correct_count",
                "value": overall["correct_count"],
                "role": "count",
                "description": "Number of correct held-out predictions.",
            },
            {
                "metric": "incorrect_count",
                "value": overall["incorrect_count"],
                "role": "count",
                "description": "Number of incorrect held-out predictions.",
            },
            {
                "metric": "correct_proportion",
                "value": overall["correct_proportion"],
                "role": "proportion",
                "description": "Proportion of correct held-out predictions.",
            },
            {
                "metric": "incorrect_proportion",
                "value": overall["incorrect_proportion"],
                "role": "proportion",
                "description": "Proportion of incorrect held-out predictions.",
            },
            {
                "metric": "test_records",
                "value": overall["test_records"],
                "role": "count",
                "description": "Total held-out statements evaluated once.",
            },
        ]
    test_metrics = pd.DataFrame(metric_rows)
    test_metrics.insert(0, "configuration_label", configuration_label)
    test_metrics.insert(0, "configuration_id", configuration_id)
    return overall, test_metrics, per_class, confusion, predictions


def model_summary_table(
    *,
    model_id: str,
    event: dict[str, Any],
    created_at: str,
    selected_configuration: str,
    cv_summaries: dict[str, dict[str, Any]],
    test_metrics: dict[str, dict[str, Any]],
    training_records: int,
    test_records: int,
    vocabulary_size: int,
    selected_count: int,
) -> pd.DataFrame:
    """Build the long-form summary sheet."""
    rows = [
        ("identity", "model_id", model_id, "Append-only model event identifier."),
        ("identity", "model_name", MODEL_NAME, "Protocol model name."),
        ("identity", "model_date", created_at[:10], "UTC execution date."),
        ("lineage", "matching_id", event["matching_id"], "Reviewed alignment event."),
        ("lineage", "split_id", event["split_id"], "Registered train/test split."),
        (
            "lineage",
            "vectorization_id",
            event["vectorization_id"],
            "Registered TF-IDF representation.",
        ),
        ("lineage", "segmentation_id", event["segmentation_id"], "Segmentation event."),
        ("lineage", "document_id", event["document_id"], "Source document."),
        ("data", "training_records", training_records, "Statements used in CV/final fitting."),
        ("data", "test_records", test_records, "Statements held out until final evaluation."),
        (
            "retention",
            "reference_vocabulary_size",
            REFERENCE_VOCABULARY_SIZE,
            "Vocabulary inspected in the released Wroblewska classifier.",
        ),
        (
            "retention",
            "reference_selected_feature_count",
            REFERENCE_SELECTED_FEATURE_COUNT,
            "Relevant n-gram count reported by Wroblewska et al.",
        ),
        (
            "retention",
            "reference_retention_proportion",
            REFERENCE_RETENTION_PROPORTION,
            "70 / 3,770, preserved proportionally rather than as an absolute count.",
        ),
        ("retention", "current_vocabulary_size", vocabulary_size, "VECT vocabulary size."),
        (
            "retention",
            "current_selected_feature_count",
            selected_count,
            "round(current vocabulary x 70 / 3,770).",
        ),
        (
            "selection",
            "feature_selection_method",
            "SelectKBest(mutual_info_classif)",
            "Refitted inside every CV training fold and independently on the complete "
            "training set for A1 and A2.",
        ),
        (
            "selection",
            "principal_selection_metric",
            "mean five-fold CV macro-F1",
            "Gives equal influence to all three classes.",
        ),
        (
            "selection",
            "A1_mean_cv_macro_f1",
            cv_summaries["A1"]["mean_macro_f1"],
            "Unweighted baseline mean.",
        ),
        (
            "selection",
            "A2_mean_cv_macro_f1",
            cv_summaries["A2"]["mean_macro_f1"],
            "Balanced adaptation mean.",
        ),
        (
            "selection",
            "selected_configuration",
            selected_configuration,
            f"Primary CV reference: {CONFIGURATIONS[selected_configuration]['label']}; "
            "the held-out report includes both A1 and A2.",
        ),
        (
            "test",
            "held_out_reporting_scope",
            "A1 and A2",
            "Both predefined configurations are reported; test scores do not select one.",
        ),
        (
            "test",
            "A1_test_macro_f1",
            test_metrics["A1"]["macro_f1"],
            "Unweighted configuration held-out macro-F1.",
        ),
        (
            "test",
            "A1_test_weighted_f1",
            test_metrics["A1"]["weighted_f1"],
            "Unweighted configuration held-out weighted F1.",
        ),
        (
            "test",
            "A1_test_accuracy",
            test_metrics["A1"]["accuracy"],
            "Unweighted configuration held-out accuracy.",
        ),
        (
            "test",
            "A2_test_macro_f1",
            test_metrics["A2"]["macro_f1"],
            "Balanced configuration held-out macro-F1.",
        ),
        (
            "test",
            "A2_test_weighted_f1",
            test_metrics["A2"]["weighted_f1"],
            "Balanced configuration held-out weighted F1.",
        ),
        (
            "test",
            "A2_test_accuracy",
            test_metrics["A2"]["accuracy"],
            "Balanced configuration held-out accuracy.",
        ),
        (
            "method",
            "model_scope",
            "proportional Wroblewska-inspired adaptation",
            "Transparent baseline; not an exact reproduction of the released pipeline.",
        ),
    ]
    return pd.DataFrame(rows, columns=["section", "item", "value", "description"])


def artifact_descriptions_table() -> pd.DataFrame:
    """Build filename, role, and description rows for all six outputs."""
    role_to_file = {
        "fitted_model": MODEL_FILENAME,
        "cross_validation_results": CV_RESULTS_FILENAME,
        "selected_features": SELECTED_FEATURES_FILENAME,
        "test_predictions": TEST_PREDICTIONS_FILENAME,
        "confusion_matrix": CONFUSION_MATRIX_FILENAME,
        "performance_report": PERFORMANCE_REPORT_FILENAME,
    }
    return pd.DataFrame(
        [
            {
                "filename": role_to_file[role],
                "role": role,
                "description": ARTIFACT_DESCRIPTIONS[role],
            }
            for role in role_to_file
        ]
    )


def excel_value(value: Any) -> Any:
    """Convert numpy/pandas scalar values to openpyxl-compatible Python values."""
    if pd.isna(value):
        return None
    if isinstance(value, np.generic):
        return value.item()
    return value


def add_dataframe_sheet(
    workbook: Workbook,
    title: str,
    data: pd.DataFrame,
    table_name: str,
) -> None:
    """Add one consistently styled, filterable worksheet."""
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(list(data.columns))
    for row in data.itertuples(index=False, name=None):
        sheet.append([excel_value(value) for value in row])

    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30

    if len(data) > 0:
        reference = f"A1:{get_column_letter(len(data.columns))}{len(data) + 1}"
        table = Table(displayName=table_name, ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    percentage_columns = {
        "macro_f1",
        "weighted_f1",
        "accuracy",
        "macro_precision",
        "macro_recall",
        "precision",
        "recall",
        "f1",
        "mutual_information_score",
    }
    for column_index, column_name in enumerate(data.columns, start=1):
        values = [str(value) for value in data[column_name].head(250) if not pd.isna(value)]
        longest = max([len(str(column_name))] + [min(len(value), 80) for value in values])
        width = min(max(longest + 2, 11), 55)
        if column_name == "candidate_text_original":
            width = 55
        elif column_name in {"description", "value"}:
            width = max(width, 30)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        if column_name in percentage_columns:
            for row_number in range(2, len(data) + 2):
                sheet.cell(row_number, column_index).number_format = "0.000000"

    wrap_columns = {
        "candidate_text_original",
        "description",
        "value",
        "ngram_text",
    }
    for column_index, column_name in enumerate(data.columns, start=1):
        if column_name in wrap_columns:
            for row_number in range(2, len(data) + 2):
                sheet.cell(row_number, column_index).alignment = Alignment(
                    vertical="top", wrap_text=True
                )

    if title == "Test_Predictions":
        correct_column = list(data.columns).index("correct") + 1
        for row_number in range(2, len(data) + 2):
            passed = bool(sheet.cell(row_number, correct_column).value)
            sheet.cell(row_number, correct_column).fill = PatternFill(
                "solid", fgColor="E2F0D9" if passed else "FCE4D6"
            )
    if title == "CV_Comparison":
        selected_column = list(data.columns).index("selected_configuration") + 1
        for row_number in range(2, len(data) + 2):
            if sheet.cell(row_number, selected_column).value:
                for cell in sheet[row_number]:
                    cell.fill = PatternFill("solid", fgColor="D9EAF7")


def write_performance_workbook(
    path: Path,
    sheets: dict[str, pd.DataFrame],
) -> None:
    """Write the eight-sheet comparison-ready performance workbook."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = MODEL_NAME
    workbook.properties.subject = "Model A cross-validation and held-out evaluation"
    workbook.properties.creator = "IG protocol project"
    table_names = {
        "Model_Summary": "TblModelSummary",
        "CV_Comparison": "TblCVComparison",
        "Test_Metrics": "TblTestMetrics",
        "Per_Class_Metrics": "TblPerClassMetrics",
        "Confusion_Matrix": "TblConfusionMatrix",
        "Selected_Features": "TblSelectedFeatures",
        "Test_Predictions": "TblTestPredictions",
        "Artifact_Descriptions": "TblArtifactDescriptions",
    }
    for title, data in sheets.items():
        add_dataframe_sheet(workbook, title, data, table_names[title])
    workbook.save(path)


def write_outputs(
    *,
    output_dir: Path,
    model_bundle: dict[str, Any],
    cv_results: pd.DataFrame,
    selected_features: pd.DataFrame,
    predictions: pd.DataFrame,
    confusion: pd.DataFrame,
    workbook_sheets: dict[str, pd.DataFrame],
) -> dict[str, Path]:
    """Write all six artifacts to a staging directory."""
    paths = {
        "fitted_model": output_dir / MODEL_FILENAME,
        "cross_validation_results": output_dir / CV_RESULTS_FILENAME,
        "selected_features": output_dir / SELECTED_FEATURES_FILENAME,
        "test_predictions": output_dir / TEST_PREDICTIONS_FILENAME,
        "confusion_matrix": output_dir / CONFUSION_MATRIX_FILENAME,
        "performance_report": output_dir / PERFORMANCE_REPORT_FILENAME,
    }
    joblib.dump(model_bundle, paths["fitted_model"], compress=3)
    cv_results.to_csv(paths["cross_validation_results"], index=False, lineterminator="\n")
    selected_features.to_csv(paths["selected_features"], index=False, lineterminator="\n")
    predictions.to_csv(paths["test_predictions"], index=False, lineterminator="\n")
    confusion.to_csv(paths["confusion_matrix"], index=False, lineterminator="\n")
    write_performance_workbook(paths["performance_report"], workbook_sheets)
    return paths


def validate_outputs(
    *,
    paths: dict[str, Path],
    test_matrix: sparse.csr_matrix,
    expected_predictions: pd.DataFrame,
    expected_selected_features: pd.DataFrame,
    expected_confusion: pd.DataFrame,
    expected_sheet_names: list[str],
    k_features: int,
) -> dict[str, Any]:
    """Read back and validate every model output artifact."""
    loaded_bundle: dict[str, Any] = joblib.load(paths["fitted_model"])
    loaded_predictions = pd.read_csv(
        paths["test_predictions"], dtype=object, keep_default_na=False
    )
    loaded_selected = pd.read_csv(
        paths["selected_features"], dtype=object, keep_default_na=False
    )
    loaded_confusion = pd.read_csv(
        paths["confusion_matrix"], dtype=object, keep_default_na=False
    )
    loaded_cv = pd.read_csv(paths["cross_validation_results"], keep_default_na=False)
    workbook = load_workbook(paths["performance_report"], read_only=True, data_only=True)
    try:
        workbook_sheets = workbook.sheetnames
    finally:
        workbook.close()

    loaded_pipelines = loaded_bundle.get("pipelines", {})
    reproduced_by_configuration: dict[str, list[str]] = {}
    expected_by_configuration: dict[str, list[str]] = {}
    for configuration_id in CONFIGURATIONS:
        if configuration_id in loaded_pipelines:
            reproduced_by_configuration[configuration_id] = (
                loaded_pipelines[configuration_id].predict(test_matrix).astype(str).tolist()
            )
        expected_by_configuration[configuration_id] = expected_predictions.loc[
            expected_predictions["configuration_id"].eq(configuration_id),
            "predicted_label",
        ].map(str).tolist()
    confusion_columns = [f"predicted_{label}" for label in CLASS_ORDER]
    loaded_confusion_values = loaded_confusion[confusion_columns].astype(int).to_numpy()
    expected_confusion_values = expected_confusion[confusion_columns].astype(int).to_numpy()
    prediction_alignment_columns = [
        "configuration_id",
        "matrix_row_index",
        "candidate_id",
    ]
    checks = {
        "all_six_output_files_exist_and_nonempty": len(paths) == 6
        and all(path.is_file() and path.stat().st_size > 0 for path in paths.values()),
        "model_bundle_contains_both_final_pipelines": set(loaded_pipelines)
        == set(CONFIGURATIONS),
        "model_bundle_class_order_preserved": tuple(loaded_bundle.get("class_order", ()))
        == CLASS_ORDER,
        "reloaded_models_reproduce_predictions": reproduced_by_configuration
        == expected_by_configuration,
        "every_configuration_predicts_every_test_statement": len(loaded_predictions)
        == len(expected_predictions)
        == len(CONFIGURATIONS) * test_matrix.shape[0],
        "prediction_configuration_ids_complete": set(
            loaded_predictions["configuration_id"].map(str)
        )
        == set(CONFIGURATIONS),
        "prediction_candidate_alignment_preserved": loaded_predictions[
            prediction_alignment_columns
        ].astype(str).values.tolist()
        == expected_predictions[prediction_alignment_columns].astype(str).values.tolist(),
        "predictions_use_only_approved_labels": set(
            loaded_predictions["predicted_label"].map(str)
        ).issubset(CLASS_ORDER),
        "selected_feature_count_exact": len(loaded_selected) == k_features,
        "selected_feature_indices_match": loaded_selected[
            "original_tfidf_feature_index"
        ].astype(int).tolist()
        == expected_selected_features["original_tfidf_feature_index"].astype(int).tolist(),
        "selected_feature_scores_match": bool(
            np.allclose(
                loaded_selected["mutual_information_score"].astype(float).to_numpy(),
                expected_selected_features["mutual_information_score"].astype(float).to_numpy(),
                rtol=0,
                atol=1e-12,
            )
        ),
        "confusion_configuration_ids_complete": set(
            loaded_confusion["configuration_id"].map(str)
        )
        == set(CONFIGURATIONS),
        "confusion_matrices_match_memory": np.array_equal(
            loaded_confusion_values, expected_confusion_values
        ),
        "each_confusion_matrix_total_matches_test_rows": all(
            int(
                loaded_confusion.loc[
                    loaded_confusion["configuration_id"].eq(configuration_id),
                    confusion_columns,
                ].astype(int).to_numpy().sum()
            )
            == test_matrix.shape[0]
            for configuration_id in CONFIGURATIONS
        ),
        "cv_contains_ten_fold_rows": int((loaded_cv["record_type"] == "fold").sum())
        == len(CONFIGURATIONS) * CV_SPLITS,
        "cv_contains_four_aggregate_rows": int(
            loaded_cv["record_type"].isin(["mean", "std"]).sum()
        )
        == len(CONFIGURATIONS) * 2,
        "workbook_has_expected_sheets_in_order": workbook_sheets == expected_sheet_names,
    }
    checks = {name: bool(value) for name, value in checks.items()}
    failed = [name for name, passed in checks.items() if not passed]
    if failed:
        raise ValueError("Model output validation failed: " + ", ".join(failed))
    return {"status": "passed", "checks": checks}


def artifact_rows(
    *,
    model_id: str,
    final_dir: Path,
    staged_paths: dict[str, Path],
    created_at: str,
) -> list[dict[str, Any]]:
    """Build model artifact registry rows from validated staged files."""
    return [
        {
            "model_id": model_id,
            "artifact_role": role,
            "file_name": path.name,
            "artifact_path": project_relative(final_dir / path.name),
            "artifact_hash_sha256": sha256_file(path),
            "size_bytes": path.stat().st_size,
            "artifact_description": ARTIFACT_DESCRIPTIONS[role],
            "created_at": created_at,
        }
        for role, path in staged_paths.items()
    ]


def register_and_publish(
    *,
    database_path: Path,
    staged_dir: Path,
    final_dir: Path,
    event_values: dict[str, Any],
    artifacts: list[dict[str, Any]],
    validation_results: dict[str, Any],
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Publish files and register the append-only event in one transaction."""
    if final_dir.exists():
        raise ValueError(f"Model output already exists: {final_dir}")
    connection = sqlite3.connect(database_path)
    published = False
    try:
        connection.execute("PRAGMA foreign_keys = ON")
        if connection.execute(
            "SELECT 1 FROM model_events WHERE model_id = ?",
            (event_values["model_id"],),
        ).fetchone():
            raise ValueError(f"Model ID is already registered: {event_values['model_id']}")
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            """
            INSERT INTO model_events (
                model_id, vectorization_id, split_id, matching_id,
                segmentation_id, document_id, model_date, created_at,
                model_name, model_designation, model_description, script_path,
                script_version, model_tool, model_tool_version,
                input_artifacts_json, input_artifacts_count, class_order_json,
                class_distributions_json, training_records_count,
                test_records_count, reference_vocabulary_size,
                reference_selected_feature_count, retention_proportion,
                retention_calculation, current_vocabulary_size,
                current_selected_feature_count, feature_selection_method,
                feature_selection_parameters_json, cross_validation_design_json,
                random_forest_configurations_json, cross_validation_results_json,
                configuration_selection_rule, selected_configuration,
                final_test_metrics_json, validation_results_json,
                output_artifacts_count, model_status, notes
            ) VALUES (
                :model_id, :vectorization_id, :split_id, :matching_id,
                :segmentation_id, :document_id, :model_date, :created_at,
                :model_name, :model_designation, :model_description, :script_path,
                :script_version, :model_tool, :model_tool_version,
                :input_artifacts_json, :input_artifacts_count, :class_order_json,
                :class_distributions_json, :training_records_count,
                :test_records_count, :reference_vocabulary_size,
                :reference_selected_feature_count, :retention_proportion,
                :retention_calculation, :current_vocabulary_size,
                :current_selected_feature_count, :feature_selection_method,
                :feature_selection_parameters_json, :cross_validation_design_json,
                :random_forest_configurations_json, :cross_validation_results_json,
                :configuration_selection_rule, :selected_configuration,
                :final_test_metrics_json, :validation_results_json,
                :output_artifacts_count, :model_status, :notes
            )
            """,
            event_values,
        )
        connection.executemany(
            """
            INSERT INTO model_artifacts (
                model_id, artifact_role, file_name, artifact_path,
                artifact_hash_sha256, size_bytes, artifact_description, created_at
            ) VALUES (
                :model_id, :artifact_role, :file_name, :artifact_path,
                :artifact_hash_sha256, :size_bytes, :artifact_description, :created_at
            )
            """,
            artifacts,
        )

        staged_dir.replace(final_dir)
        published = True
        for artifact in artifacts:
            path = resolve_project_path(artifact["artifact_path"])
            if sha256_file(path) != artifact["artifact_hash_sha256"]:
                raise RuntimeError(f"Published hash mismatch: {artifact['file_name']}")
            if path.stat().st_size != artifact["size_bytes"]:
                raise RuntimeError(f"Published size mismatch: {artifact['file_name']}")

        artifact_count = connection.execute(
            "SELECT COUNT(*) FROM model_artifacts WHERE model_id = ?",
            (event_values["model_id"],),
        ).fetchone()[0]
        if artifact_count != len(artifacts):
            raise RuntimeError("Model artifact registry count is incorrect.")
        if connection.execute("PRAGMA foreign_key_check").fetchall():
            raise RuntimeError("Foreign-key validation failed during model registration.")
        integrity = connection.execute("PRAGMA integrity_check").fetchall()
        if integrity != [("ok",)]:
            raise RuntimeError(f"SQLite integrity check failed: {integrity}")

        final_validation = {
            **validation_results,
            "published_artifact_hashes_match_registry": True,
            "database_event_corresponds_to_exported_artifacts": True,
            "database_registration_transactional": True,
            "database_foreign_key_check_passed": True,
            "database_integrity_check_passed": True,
        }
        connection.execute(
            "UPDATE model_events SET validation_results_json = ? WHERE model_id = ?",
            (json_text(final_validation), event_values["model_id"]),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        if published and final_dir.exists() and not staged_dir.exists():
            final_dir.replace(staged_dir)
        raise
    finally:
        connection.close()

    with sqlite3.connect(database_path) as verification:
        event_count = verification.execute(
            "SELECT COUNT(*) FROM model_events WHERE model_id = ?",
            (event_values["model_id"],),
        ).fetchone()[0]
        artifact_count = verification.execute(
            "SELECT COUNT(*) FROM model_artifacts WHERE model_id = ?",
            (event_values["model_id"],),
        ).fetchone()[0]
    if event_count != 1 or artifact_count != len(artifacts):
        raise RuntimeError("Committed model registration could not be verified.")
    return (
        {
            "status": "registered",
            "model_id": event_values["model_id"],
            "artifact_records": artifact_count,
        },
        final_validation,
    )


def run_model_a(
    *,
    database_path: Path,
    output_base_dir: Path,
    vectorization_id: str | None = None,
) -> dict[str, Any]:
    """Validate, train, evaluate, export, publish, and register Model A."""
    created_at = utc_now_iso()
    ensure_registry_schema(database_path)
    vector_event, vector_artifacts = load_vectorization_event(
        database_path, vectorization_id
    )
    inputs = load_and_validate_inputs(vector_event, vector_artifacts)
    train_matrix: sparse.csr_matrix = inputs["train_matrix"]
    test_matrix: sparse.csr_matrix = inputs["test_matrix"]
    train_index: pd.DataFrame = inputs["train_index"]
    test_index: pd.DataFrame = inputs["test_index"]
    features: pd.DataFrame = inputs["features"]

    vocabulary_size = train_matrix.shape[1]
    k_features = selected_feature_count(vocabulary_size)
    if vocabulary_size != EXPECTED_CURRENT_VOCABULARY_SIZE:
        raise ValueError(
            "VECT vocabulary differs from the approved Model A input: "
            f"expected {EXPECTED_CURRENT_VOCABULARY_SIZE:,}, found {vocabulary_size:,}."
        )
    if k_features != EXPECTED_SELECTED_FEATURE_COUNT or k_features != 436:
        raise ValueError(f"Feature-retention calculation did not produce 436: {k_features}.")

    labels = train_index["reviewer_gold_label"].map(str).to_numpy()
    folds, fold_validation = build_cv_folds(labels)
    cv_results, cv_summaries = run_cross_validation(
        train_matrix, labels, folds, k_features
    )
    selected_configuration = select_configuration(cv_summaries)
    cv_results.loc[
        cv_results["configuration_id"].eq(selected_configuration),
        "selected_configuration",
    ] = True

    final_pipelines: dict[str, Pipeline] = {}
    selected_feature_tables: dict[str, pd.DataFrame] = {}
    test_results: dict[str, dict[str, Any]] = {}
    test_metric_tables: list[pd.DataFrame] = []
    per_class_tables: list[pd.DataFrame] = []
    confusion_tables: list[pd.DataFrame] = []
    prediction_tables: list[pd.DataFrame] = []
    for configuration_id, configuration in CONFIGURATIONS.items():
        pipeline = build_pipeline(configuration["class_weight"], k_features)
        pipeline.fit(train_matrix, labels)
        selected_feature_tables[configuration_id] = build_selected_features(
            pipeline, features
        )
        (
            overall,
            configuration_test_metrics,
            configuration_per_class,
            configuration_confusion,
            configuration_predictions,
        ) = build_evaluation_outputs(
            configuration_id,
            pipeline,
            test_matrix,
            test_index,
        )
        test_results[configuration_id] = overall
        test_metric_tables.append(configuration_test_metrics)
        per_class_tables.append(configuration_per_class)
        confusion_tables.append(configuration_confusion)
        prediction_tables.append(configuration_predictions)
        # A standard-library partial keeps the bundle loadable outside __main__.
        pipeline.named_steps["feature_selection"].score_func = partial(
            mutual_info_classif, random_state=RANDOM_STATE
        )
        final_pipelines[configuration_id] = pipeline

    selected_features = selected_feature_tables["A1"]
    test_metrics = pd.concat(test_metric_tables, ignore_index=True)
    per_class = pd.concat(per_class_tables, ignore_index=True)
    confusion = pd.concat(confusion_tables, ignore_index=True)
    predictions = pd.concat(prediction_tables, ignore_index=True)
    if len(selected_features) != k_features:
        raise RuntimeError(f"Final models selected {len(selected_features)}, not {k_features}.")
    selected_feature_tables_identical = selected_feature_tables["A1"].equals(
        selected_feature_tables["A2"]
    )
    if not selected_feature_tables_identical:
        raise RuntimeError("A1 and A2 final feature-selection tables differ unexpectedly.")

    model_id = next_model_id(database_path)
    output_base_dir.mkdir(parents=True, exist_ok=True)
    final_dir = output_base_dir / model_id
    staged_dir = Path(
        tempfile.mkdtemp(prefix=f".{model_id}_", dir=output_base_dir)
    )
    model_bundle = {
        "artifact_type": "paired_model_a_configuration_bundle",
        "model_id": model_id,
        "script_version": SCRIPT_VERSION,
        "class_order": list(CLASS_ORDER),
        "cv_selected_primary_configuration": selected_configuration,
        "held_out_reported_configurations": list(CONFIGURATIONS),
        "test_scores_used_for_selection": False,
        "pipelines": final_pipelines,
    }

    summary = model_summary_table(
        model_id=model_id,
        event=vector_event,
        created_at=created_at,
        selected_configuration=selected_configuration,
        cv_summaries=cv_summaries,
        test_metrics=test_results,
        training_records=len(train_index),
        test_records=len(test_index),
        vocabulary_size=vocabulary_size,
        selected_count=k_features,
    )
    workbook_sheets = {
        "Model_Summary": summary,
        "CV_Comparison": cv_results,
        "Test_Metrics": test_metrics,
        "Per_Class_Metrics": per_class,
        "Confusion_Matrix": confusion,
        "Selected_Features": selected_features,
        "Test_Predictions": predictions,
        "Artifact_Descriptions": artifact_descriptions_table(),
    }
    expected_sheet_names = list(workbook_sheets)

    try:
        staged_paths = write_outputs(
            output_dir=staged_dir,
            model_bundle=model_bundle,
            cv_results=cv_results,
            selected_features=selected_features,
            predictions=predictions,
            confusion=confusion,
            workbook_sheets=workbook_sheets,
        )
        output_validation = validate_outputs(
            paths=staged_paths,
            test_matrix=test_matrix,
            expected_predictions=predictions,
            expected_selected_features=selected_features,
            expected_confusion=confusion,
            expected_sheet_names=expected_sheet_names,
            k_features=k_features,
        )
        additional_checks = {
            "exactly_436_features_selected": len(selected_features) == 436,
            "A1_and_A2_final_selected_features_identical": selected_feature_tables_identical,
            "selected_indices_map_to_feature_table": selected_features[
                "ngram_text"
            ].notna().all(),
            "both_configurations_fitted_on_complete_training_partition": set(
                final_pipelines
            )
            == set(CONFIGURATIONS),
            "test_prediction_count_equals_two_times_test_rows": len(predictions)
            == len(CONFIGURATIONS) * len(test_index),
            "each_per_class_support_total_equals_test_rows": all(
                int(
                    per_class.loc[
                        per_class["configuration_id"].eq(configuration_id), "support"
                    ].sum()
                )
                == len(test_index)
                for configuration_id in CONFIGURATIONS
            ),
            "each_metric_count_total_equals_test_rows": all(
                test_results[configuration_id]["correct_count"]
                + test_results[configuration_id]["incorrect_count"]
                == len(test_index)
                for configuration_id in CONFIGURATIONS
            ),
            "test_labels_not_used_for_selection": True,
            "test_scores_not_used_to_choose_primary_configuration": True,
            "feature_selection_refitted_inside_each_cv_pipeline": True,
            "identical_materialized_folds_used_for_A1_and_A2": True,
            "both_predefined_configurations_reported_on_held_out_partition": True,
        }
        additional_checks = {
            name: bool(passed) for name, passed in additional_checks.items()
        }
        failed = [name for name, passed in additional_checks.items() if not passed]
        if failed:
            raise RuntimeError("Final correctness validation failed: " + ", ".join(failed))
        validation_results = {
            "status": "passed",
            "input_validation": inputs["validation"],
            "cross_validation_fold_validation": fold_validation,
            "output_validation": output_validation,
            "final_correctness_checks": additional_checks,
        }

        artifacts = artifact_rows(
            model_id=model_id,
            final_dir=final_dir,
            staged_paths=staged_paths,
            created_at=created_at,
        )
        input_artifact_records = [
            {
                "artifact_role": role,
                "file_name": values["file_name"],
                "artifact_path": values["artifact_path"],
                "artifact_hash_sha256": values["artifact_hash_sha256"],
                "size_bytes": int(values["size_bytes"]),
            }
            for role, values in sorted(vector_artifacts.items())
        ]
        configuration_records = {
            configuration_id: {
                **RANDOM_FOREST_PARAMETERS,
                "class_weight": configuration["class_weight"],
                "configuration_label": configuration["label"],
            }
            for configuration_id, configuration in CONFIGURATIONS.items()
        }
        cv_registry_results = {
            "stored_numeric_precision_decimal_places": STORED_METRIC_DECIMALS,
            "configurations": cv_summaries,
            "records": cv_results.to_dict(orient="records"),
        }
        final_metrics_registry = {
            "class_order": list(CLASS_ORDER),
            "held_out_reported_configurations": list(CONFIGURATIONS),
            "cv_selected_primary_configuration": selected_configuration,
            "test_scores_used_for_selection": False,
            "configurations": {
                configuration_id: {
                    "configuration_label": CONFIGURATIONS[configuration_id]["label"],
                    "overall": test_results[configuration_id],
                    "per_class": per_class.loc[
                        per_class["configuration_id"].eq(configuration_id)
                    ].to_dict(orient="records"),
                    "confusion_matrix": {
                        "rows_are_actual": True,
                        "columns_are_predicted": True,
                        "values": confusion.loc[
                            confusion["configuration_id"].eq(configuration_id),
                            [f"predicted_{label}" for label in CLASS_ORDER],
                        ].astype(int).values.tolist(),
                    },
                }
                for configuration_id in CONFIGURATIONS
            },
        }
        event_values = {
            "model_id": model_id,
            "vectorization_id": vector_event["vectorization_id"],
            "split_id": vector_event["split_id"],
            "matching_id": vector_event["matching_id"],
            "segmentation_id": vector_event["segmentation_id"],
            "document_id": vector_event["document_id"],
            "model_date": created_at[:10],
            "created_at": created_at,
            "model_name": MODEL_NAME,
            "model_designation": MODEL_DESIGNATION,
            "model_description": MODEL_DESCRIPTION,
            "script_path": project_relative(Path(__file__)),
            "script_version": SCRIPT_VERSION,
            "model_tool": "scikit-learn Pipeline, SelectKBest, and RandomForestClassifier",
            "model_tool_version": sklearn.__version__,
            "input_artifacts_json": json_text(input_artifact_records),
            "input_artifacts_count": len(input_artifact_records),
            "class_order_json": json_text(list(CLASS_ORDER)),
            "class_distributions_json": json_text(inputs["class_distributions"]),
            "training_records_count": len(train_index),
            "test_records_count": len(test_index),
            "reference_vocabulary_size": REFERENCE_VOCABULARY_SIZE,
            "reference_selected_feature_count": REFERENCE_SELECTED_FEATURE_COUNT,
            "retention_proportion": REFERENCE_RETENTION_PROPORTION,
            "retention_calculation": (
                f"round({vocabulary_size} * {REFERENCE_SELECTED_FEATURE_COUNT} / "
                f"{REFERENCE_VOCABULARY_SIZE}) = {k_features}"
            ),
            "current_vocabulary_size": vocabulary_size,
            "current_selected_feature_count": k_features,
            "feature_selection_method": (
                "SelectKBest with reproducible mutual_info_classif, fitted inside "
                "each CV fold and independently for A1 and A2 on all training records"
            ),
            "feature_selection_parameters_json": json_text(
                {
                    "k": k_features,
                    "score_func": "mutual_info_classif",
                    "random_state": RANDOM_STATE,
                }
            ),
            "cross_validation_design_json": json_text(
                {
                    "type": "StratifiedKFold",
                    "n_splits": CV_SPLITS,
                    "shuffle": True,
                    "random_state": RANDOM_STATE,
                    "same_folds_for_all_configurations": True,
                    "principal_metric": "macro_f1",
                    "held_out_reported_configurations": list(CONFIGURATIONS),
                    "test_scores_used_for_selection": False,
                    "fold_class_distributions": fold_validation[
                        "fold_class_distributions"
                    ],
                }
            ),
            "random_forest_configurations_json": json_text(configuration_records),
            "cross_validation_results_json": json_text(cv_registry_results),
            "configuration_selection_rule": (
                "Select the greatest mean five-fold CV macro-F1 at 12 decimal places; "
                "select A1 on a stored-precision tie. Retain that result as the primary "
                "CV reference, report held-out metrics for both predefined configurations, "
                "and never use test performance to change the selection."
            ),
            "selected_configuration": selected_configuration,
            "final_test_metrics_json": json_text(final_metrics_registry),
            "validation_results_json": json_text(validation_results),
            "output_artifacts_count": len(artifacts),
            "model_status": "Completed",
            "notes": (
                "Protocol revision 1.1 reports held-out results for both A1 and A2 so "
                "the later B1/B2 full-feature models can use the same comparison structure. "
                "MODEL_000001 and its winner-only test report remain unchanged. Because the "
                "A1 test result was observed before this expanded comparison was requested, "
                "the paired held-out comparison is post-hoc relative to the initial run; "
                "test scores are descriptive and do not select or tune a configuration. "
                "Proportional Wroblewska-inspired adaptation, not an exact reproduction. "
                "The reference does not explicitly justify the absolute count 70, so its "
                "vocabulary proportion is preserved. The released pipeline also used "
                "Random-Forest recursive feature elimination, which this transparent "
                "baseline omits. With sparse TF-IDF input, scikit-learn's mutual-information "
                "implementation treats feature values as discrete; this is a practical "
                "ranking approximation for continuous TF-IDF weights."
            ),
        }
        registration, final_validation = register_and_publish(
            database_path=database_path,
            staged_dir=staged_dir,
            final_dir=final_dir,
            event_values=event_values,
            artifacts=artifacts,
            validation_results=validation_results,
        )
    finally:
        if staged_dir.exists():
            shutil.rmtree(staged_dir)

    print(f"linked_vectorization_id: {vector_event['vectorization_id']}")
    print(f"training_records: {len(train_index)}")
    print(f"test_records: {len(test_index)}")
    print(
        "feature_retention: "
        f"round({vocabulary_size:,} x {REFERENCE_SELECTED_FEATURE_COUNT} / "
        f"{REFERENCE_VOCABULARY_SIZE:,}) = {k_features}"
    )
    print(f"A1_mean_cv_macro_f1: {cv_summaries['A1']['mean_macro_f1']:.6f}")
    print(f"A2_mean_cv_macro_f1: {cv_summaries['A2']['mean_macro_f1']:.6f}")
    print(f"cv_selected_primary_configuration: {selected_configuration}")
    for configuration_id in CONFIGURATIONS:
        print(
            f"{configuration_id}_test_metrics: "
            f"macro_f1={test_results[configuration_id]['macro_f1']:.6f}, "
            f"weighted_f1={test_results[configuration_id]['weighted_f1']:.6f}, "
            f"accuracy={test_results[configuration_id]['accuracy']:.6f}"
        )
    print(f"output_directory: {project_relative(final_dir)}")
    print(f"registered_model_id: {model_id}")
    print("final_validation_status: passed")

    return {
        "model_id": model_id,
        "vectorization_id": vector_event["vectorization_id"],
        "split_id": vector_event["split_id"],
        "matching_id": vector_event["matching_id"],
        "output_directory": project_relative(final_dir),
        "training_records": len(train_index),
        "test_records": len(test_index),
        "vocabulary_size": vocabulary_size,
        "selected_feature_count": k_features,
        "retention_proportion": REFERENCE_RETENTION_PROPORTION,
        "cv_summaries": cv_summaries,
        "selected_configuration": selected_configuration,
        "test_metrics": test_results,
        "per_class_metrics": per_class.to_dict(orient="records"),
        "registration": registration,
        "validation": final_validation,
        "artifacts": artifacts,
    }


def parse_arguments() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(
        description="Train and register proportional Random Forest Model A."
    )
    parser.add_argument(
        "--database", default=str(DATABASE_PATH), help="Corpus inventory SQLite database."
    )
    parser.add_argument(
        "--vectorization-id",
        default=None,
        help="Registered vectorization ID. Default: latest completed vectorization.",
    )
    parser.add_argument(
        "--output-dir",
        default=str(DEFAULT_OUTPUT_DIR),
        help="Base directory for append-only MODEL_###### output folders.",
    )
    return parser.parse_args()


def main() -> None:
    """Command-line entry point."""
    args = parse_arguments()
    try:
        run_model_a(
            database_path=resolve_project_path(args.database),
            output_base_dir=resolve_project_path(args.output_dir),
            vectorization_id=args.vectorization_id,
        )
    except Exception as error:
        raise SystemExit(f"Model A failed: {error}") from error


if __name__ == "__main__":
    main()
