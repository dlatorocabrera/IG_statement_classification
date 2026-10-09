#!/usr/bin/env python3
"""Train and register paired full-feature Random Forest Model B."""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sqlite3
import sys
import tempfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd
import sklearn
from openpyxl import Workbook, load_workbook
from scipy import sparse
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import confusion_matrix, precision_recall_fscore_support
from sklearn.pipeline import Pipeline


SCRIPT_DIR = Path(__file__).resolve().parent
BASE_SCRIPT_PATH = SCRIPT_DIR / "09_train_random_forest_model_a.py"
BASE_SPEC = importlib.util.spec_from_file_location(
    "task3_random_forest_model_a_base", BASE_SCRIPT_PATH
)
if BASE_SPEC is None or BASE_SPEC.loader is None:
    raise RuntimeError(f"Could not load shared Model A utilities: {BASE_SCRIPT_PATH}")
BASE = importlib.util.module_from_spec(BASE_SPEC)
sys.modules[BASE_SPEC.name] = BASE
BASE_SPEC.loader.exec_module(BASE)


DATABASE_PATH = BASE.DATABASE_PATH
DEFAULT_OUTPUT_DIR = BASE.DEFAULT_OUTPUT_DIR
SCRIPT_NAME = "10_train_random_forest_model_b.py"
SCRIPT_VERSION = "1.0"
MODEL_NAME = "Model B: paired full-feature Random Forests"
MODEL_DESIGNATION = "Model B"
MODEL_DESCRIPTION = (
    "Paired three-class Random Forest evaluation comparing unweighted B1 and "
    "class-balanced B2 while retaining all registered TF-IDF features."
)

RANDOM_STATE = BASE.RANDOM_STATE
CLASS_ORDER = BASE.CLASS_ORDER
CV_SPLITS = BASE.CV_SPLITS
STORED_METRIC_DECIMALS = BASE.STORED_METRIC_DECIMALS
EXPECTED_VOCABULARY_SIZE = BASE.EXPECTED_CURRENT_VOCABULARY_SIZE
RANDOM_FOREST_PARAMETERS = BASE.RANDOM_FOREST_PARAMETERS

CONFIGURATIONS = {
    "B1": {"class_weight": None, "label": "full-feature unweighted baseline"},
    "B2": {"class_weight": "balanced", "label": "full-feature balanced adaptation"},
}

MODEL_FILENAME = "model_b_full_features_rf.joblib"
CV_RESULTS_FILENAME = "model_b_cv_results.csv"
FEATURES_FILENAME = "model_b_all_features.csv"
TEST_PREDICTIONS_FILENAME = "model_b_test_predictions.csv"
CONFUSION_MATRIX_FILENAME = "model_b_confusion_matrix.csv"
PERFORMANCE_REPORT_FILENAME = "model_b_full_features_rf_performance.xlsx"

ARTIFACT_DESCRIPTIONS = {
    "fitted_model": (
        "Fitted-model bundle containing the final full-feature B1 and B2 pipelines, "
        "configuration metadata, class order, and CV-selected primary reference."
    ),
    "cross_validation_results": (
        "Combined fold-level and aggregate comparison of A1, A2, B1, and B2; "
        "registered A results are reused and B results are computed in this event."
    ),
    "selected_features": (
        "Complete registered TF-IDF feature table showing that all 23,464 columns "
        "are included in both B1 and B2 without supervised selection."
    ),
    "test_predictions": (
        "Long-form candidate-aligned held-out predictions and correctness flags for "
        "both B1 and B2."
    ),
    "confusion_matrix": (
        "Held-out confusion matrices for B1 and B2 using the predefined class order."
    ),
    "performance_report": (
        "Human-readable workbook containing lineage, four-configuration CV comparison, "
        "paired Model B test performance, all features, predictions, and artifact notes."
    ),
}


def utc_now_iso() -> str:
    """Return a UTC timestamp without microseconds."""
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def ensure_general_model_registry_schema(database_path: Path) -> dict[str, Any]:
    """Generalize the configuration constraint while preserving every registry row."""
    BASE.ensure_registry_schema(database_path)
    with sqlite3.connect(database_path) as connection:
        table_sql_row = connection.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='model_events'"
        ).fetchone()
    if table_sql_row is None:
        raise RuntimeError("model_events was not created.")
    table_sql = str(table_sql_row[0])
    restrictive = "selected_configuration IN ('A1', 'A2')" in table_sql
    generalized = "length(trim(selected_configuration)) > 0" in table_sql
    if generalized:
        return {"status": "already_general", "rows_preserved": True}
    if not restrictive:
        raise RuntimeError("Unrecognized model_events selected-configuration constraint.")

    event_ddl = BASE.MODEL_REGISTRY_SCHEMA.split(
        "CREATE TABLE IF NOT EXISTS model_artifacts", maxsplit=1
    )[0].strip()
    event_ddl = event_ddl.replace(
        "CREATE TABLE IF NOT EXISTS model_events",
        "CREATE TABLE model_events_new",
        1,
    ).replace(
        "CHECK (selected_configuration IN ('A1', 'A2'))",
        "CHECK (length(trim(selected_configuration)) > 0)",
        1,
    )

    connection = sqlite3.connect(database_path)
    try:
        connection.execute("PRAGMA foreign_keys = OFF")
        before_rows = connection.execute("SELECT COUNT(*) FROM model_events").fetchone()[0]
        columns = [
            str(row[1]) for row in connection.execute("PRAGMA table_info(model_events)")
        ]
        quoted_columns = ", ".join(f'"{column}"' for column in columns)
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(event_ddl)
        connection.execute(
            f"INSERT INTO model_events_new ({quoted_columns}) "
            f"SELECT {quoted_columns} FROM model_events"
        )
        copied_rows = connection.execute(
            "SELECT COUNT(*) FROM model_events_new"
        ).fetchone()[0]
        if copied_rows != before_rows:
            raise RuntimeError("Model registry migration did not preserve every event row.")
        connection.execute("DROP TABLE model_events")
        connection.execute("ALTER TABLE model_events_new RENAME TO model_events")
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_model_events_vectorization_created
            ON model_events(vectorization_id, created_at)
            """
        )
        connection.commit()
        connection.execute("PRAGMA foreign_keys = ON")
        if connection.execute("PRAGMA foreign_key_check").fetchall():
            raise RuntimeError("Foreign-key check failed after model registry migration.")
        if connection.execute("PRAGMA integrity_check").fetchall() != [("ok",)]:
            raise RuntimeError("SQLite integrity check failed after registry migration.")
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    return {
        "status": "migrated",
        "previous_constraint": "A1/A2 only",
        "current_constraint": "any nonblank configuration identifier",
        "preserved_model_event_rows": int(before_rows),
        "rows_preserved": True,
    }


def build_pipeline(class_weight: str | None) -> Pipeline:
    """Build a full-feature Random Forest pipeline with no selection step."""
    return Pipeline(
        [
            (
                "classifier",
                RandomForestClassifier(
                    **RANDOM_FOREST_PARAMETERS,
                    class_weight=class_weight,
                ),
            )
        ]
    )


def run_cross_validation(
    matrix: sparse.csr_matrix,
    labels: np.ndarray,
    folds: list[tuple[np.ndarray, np.ndarray]],
) -> tuple[pd.DataFrame, dict[str, dict[str, Any]]]:
    """Evaluate B1 and B2 on the same folds previously used by Model A."""
    metric_names = (
        "macro_f1",
        "weighted_f1",
        "accuracy",
        "macro_precision",
        "macro_recall",
    )
    rows: list[dict[str, Any]] = []
    summaries: dict[str, dict[str, Any]] = {}
    for configuration_id, configuration in CONFIGURATIONS.items():
        fold_rows: list[dict[str, Any]] = []
        for fold_number, (training_rows, validation_rows) in enumerate(folds, start=1):
            pipeline = build_pipeline(configuration["class_weight"])
            pipeline.fit(matrix[training_rows], labels[training_rows])
            classifier = pipeline.named_steps["classifier"]
            if int(classifier.n_features_in_) != matrix.shape[1]:
                raise RuntimeError(
                    f"{configuration_id} fold {fold_number} did not use every feature."
                )
            predicted = pipeline.predict(matrix[validation_rows])
            metrics = BASE.calculate_metrics(labels[validation_rows], predicted)
            row = {
                "configuration_id": configuration_id,
                "configuration_label": configuration["label"],
                "class_weight": configuration["class_weight"] or "None",
                "record_type": "fold",
                "fold": fold_number,
                "training_records": len(training_rows),
                "validation_records": len(validation_rows),
                **metrics,
                "selected_configuration": False,
            }
            rows.append(row)
            fold_rows.append(row)
        summary: dict[str, Any] = {}
        for metric in metric_names:
            values = np.asarray([row[metric] for row in fold_rows], dtype=float)
            summary[f"mean_{metric}"] = float(values.mean())
            summary[f"std_{metric}"] = float(values.std(ddof=1))
        summaries[configuration_id] = summary
        for record_type in ("mean", "std"):
            rows.append(
                {
                    "configuration_id": configuration_id,
                    "configuration_label": configuration["label"],
                    "class_weight": configuration["class_weight"] or "None",
                    "record_type": record_type,
                    "fold": "all",
                    "training_records": "",
                    "validation_records": "",
                    **{
                        metric: summary[f"{record_type}_{metric}"]
                        for metric in metric_names
                    },
                    "selected_configuration": False,
                }
            )
    return pd.DataFrame(rows), summaries


def select_configuration(
    summaries: dict[str, dict[str, Any]],
    decimal_places: int = STORED_METRIC_DECIMALS,
) -> str:
    """Choose the greatest mean CV macro-F1, selecting B1 on a precision tie."""
    b1 = round(float(summaries["B1"]["mean_macro_f1"]), decimal_places)
    b2 = round(float(summaries["B2"]["mean_macro_f1"]), decimal_places)
    return "B1" if b1 >= b2 else "B2"


def build_all_features(features: pd.DataFrame) -> pd.DataFrame:
    """Build an auditable table showing that every TF-IDF column is included."""
    output = features.copy()
    output["feature_index"] = output["feature_index"].astype(int)
    output["ngram_length"] = output["ngram_length"].astype(int)
    output["idf_value"] = output["idf_value"].astype(float)
    output = output.rename(columns={"feature_index": "original_tfidf_feature_index"})
    output.insert(0, "model_feature_position", np.arange(1, len(output) + 1))
    output.insert(4, "feature_inclusion_method", "all_features")
    output.insert(5, "mutual_information_score", "")
    return output.loc[
        :,
        [
            "model_feature_position",
            "original_tfidf_feature_index",
            "ngram_text",
            "ngram_length",
            "feature_inclusion_method",
            "mutual_information_score",
            "idf_value",
        ],
    ]


def build_evaluation_outputs(
    configuration_id: str,
    pipeline: Pipeline,
    matrix: sparse.csr_matrix,
    test_index: pd.DataFrame,
) -> tuple[dict[str, Any], pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Evaluate one Model B configuration and construct long-form report tables."""
    configuration = CONFIGURATIONS[configuration_id]
    actual = test_index["reviewer_gold_label"].map(str).to_numpy()
    predicted = pipeline.predict(matrix).astype(str)
    overall = BASE.calculate_metrics(actual, predicted)
    correct = actual == predicted
    overall.update(
        {
            "correct_count": int(correct.sum()),
            "incorrect_count": int((~correct).sum()),
            "correct_proportion": float(correct.mean()),
            "incorrect_proportion": float((~correct).mean()),
            "test_records": int(len(actual)),
        }
    )
    precision, recall, f1, support = precision_recall_fscore_support(
        actual, predicted, labels=CLASS_ORDER, zero_division=0
    )
    per_class = pd.DataFrame(
        {
            "configuration_id": configuration_id,
            "configuration_label": configuration["label"],
            "class": CLASS_ORDER,
            "precision": precision,
            "recall": recall,
            "f1": f1,
            "support": support.astype(int),
        }
    )
    values = confusion_matrix(actual, predicted, labels=CLASS_ORDER)
    confusion = pd.DataFrame(
        values, columns=[f"predicted_{label}" for label in CLASS_ORDER]
    )
    confusion.insert(0, "configuration_label", configuration["label"])
    confusion.insert(0, "configuration_id", configuration_id)
    confusion.insert(2, "actual_label", CLASS_ORDER)
    confusion["actual_total"] = values.sum(axis=1)
    predictions = pd.DataFrame(
        {
            "configuration_id": configuration_id,
            "configuration_label": configuration["label"],
            "matrix_row_index": test_index["matrix_row_index"].astype(int),
            "candidate_id": test_index["candidate_id"].map(str),
            "candidate_text_original": test_index["candidate_text_original"].map(str),
            "gold_label": actual,
            "predicted_label": predicted,
            "correct": correct,
        }
    )
    descriptions = {
        "macro_f1": ("principal metric", "Unweighted mean class-specific F1."),
        "weighted_f1": ("secondary metric", "Support-weighted mean class-specific F1."),
        "accuracy": ("secondary metric", "Proportion classified correctly."),
        "macro_precision": ("secondary metric", "Unweighted mean class precision."),
        "macro_recall": ("secondary metric", "Unweighted mean class recall."),
        "correct_count": ("count", "Number of correct held-out predictions."),
        "incorrect_count": ("count", "Number of incorrect held-out predictions."),
        "correct_proportion": ("proportion", "Proportion of correct predictions."),
        "incorrect_proportion": ("proportion", "Proportion of incorrect predictions."),
        "test_records": ("count", "Total held-out statements evaluated."),
    }
    test_metrics = pd.DataFrame(
        [
            {
                "configuration_id": configuration_id,
                "configuration_label": configuration["label"],
                "metric": metric,
                "value": overall[metric],
                "role": descriptions[metric][0],
                "description": descriptions[metric][1],
            }
            for metric in descriptions
        ]
    )
    return overall, test_metrics, per_class, confusion, predictions


def load_model_a_cv_reference(
    database_path: Path,
    vectorization_id: str,
) -> dict[str, Any]:
    """Load the latest paired A1/A2 CV event for the same vectorization."""
    with sqlite3.connect(database_path) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            """
            SELECT * FROM model_events
            WHERE vectorization_id = ?
              AND model_designation = 'Model A'
              AND model_status = 'Completed'
            ORDER BY created_at DESC, model_id DESC
            """,
            (vectorization_id,),
        ).fetchall()
    for row in rows:
        event = dict(row)
        metrics = json.loads(event["final_test_metrics_json"])
        cv = json.loads(event["cross_validation_results_json"])
        if set(metrics.get("held_out_reported_configurations", [])) == {"A1", "A2"}:
            return {
                "model_id": event["model_id"],
                "event": event,
                "records": pd.DataFrame(cv["records"]),
                "summaries": cv["configurations"],
                "cv_design": json.loads(event["cross_validation_design_json"]),
            }
    raise ValueError(
        f"No paired A1/A2 model event was found for vectorization {vectorization_id}."
    )


def build_four_model_cv_comparison(
    *,
    model_a_reference: dict[str, Any],
    model_b_id: str,
    model_b_results: pd.DataFrame,
    model_b_selected: str,
) -> pd.DataFrame:
    """Combine authoritative A results with the newly computed B results."""
    a = model_a_reference["records"].copy()
    b = model_b_results.copy()
    for frame, family, scope, source_id in (
        (a, "A", "436 mutual-information-selected features", model_a_reference["model_id"]),
        (b, "B", "all 23,464 TF-IDF features", model_b_id),
    ):
        frame.insert(0, "source_model_event_id", source_id)
        frame.insert(0, "feature_scope", scope)
        frame.insert(0, "model_family", family)
    b.loc[
        b["configuration_id"].eq(model_b_selected),
        "selected_configuration",
    ] = True
    combined = pd.concat([a, b], ignore_index=True)
    expected = {"A1", "A2", "B1", "B2"}
    if set(combined["configuration_id"]) != expected:
        raise RuntimeError("Four-model CV comparison does not contain A1/A2/B1/B2.")
    fold_rows = combined.loc[combined["record_type"].eq("fold")]
    for fold_number in range(1, CV_SPLITS + 1):
        fold = fold_rows.loc[fold_rows["fold"].astype(int).eq(fold_number)]
        if len(fold) != 4:
            raise RuntimeError(f"Four-model comparison fold {fold_number} is incomplete.")
        if fold["training_records"].astype(int).nunique() != 1:
            raise RuntimeError(f"Training counts differ in fold {fold_number}.")
        if fold["validation_records"].astype(int).nunique() != 1:
            raise RuntimeError(f"Validation counts differ in fold {fold_number}.")
    return combined


def model_summary_table(
    *,
    model_id: str,
    vector_event: dict[str, Any],
    model_a_reference: dict[str, Any],
    created_at: str,
    cv_summaries: dict[str, dict[str, Any]],
    selected_configuration: str,
    test_results: dict[str, dict[str, Any]],
    training_records: int,
    test_records: int,
    vocabulary_size: int,
) -> pd.DataFrame:
    """Build Model B's long-form summary sheet."""
    a_summaries = model_a_reference["summaries"]
    rows = [
        ("identity", "model_id", model_id, "Append-only Model B event."),
        ("identity", "model_name", MODEL_NAME, "Protocol model name."),
        ("identity", "model_date", created_at[:10], "UTC execution date."),
        ("lineage", "matching_id", vector_event["matching_id"], "Reviewed alignment."),
        ("lineage", "split_id", vector_event["split_id"], "Registered split."),
        ("lineage", "vectorization_id", vector_event["vectorization_id"], "TF-IDF event."),
        ("comparison", "model_a_reference_id", model_a_reference["model_id"], "Registered paired A results reused for comparison."),
        ("data", "training_records", training_records, "Complete training partition."),
        ("data", "test_records", test_records, "Held-out reporting partition."),
        ("features", "available_features", vocabulary_size, "Registered TF-IDF columns."),
        ("features", "used_features", vocabulary_size, "All columns used by B1 and B2."),
        ("features", "feature_selection", "none", "No supervised feature selection."),
        ("cv", "A1_mean_macro_f1", a_summaries["A1"]["mean_macro_f1"], "Registered Model A result."),
        ("cv", "A2_mean_macro_f1", a_summaries["A2"]["mean_macro_f1"], "Registered Model A result."),
        ("cv", "B1_mean_macro_f1", cv_summaries["B1"]["mean_macro_f1"], "Full-feature unweighted result."),
        ("cv", "B2_mean_macro_f1", cv_summaries["B2"]["mean_macro_f1"], "Full-feature balanced result."),
        ("cv", "model_b_primary_reference", selected_configuration, "Selected within Model B by mean CV macro-F1 only."),
        ("test", "held_out_reporting_scope", "B1 and B2", "Both predefined B configurations reported without test-based selection."),
    ]
    for configuration_id in CONFIGURATIONS:
        for metric in ("macro_f1", "weighted_f1", "accuracy"):
            rows.append(
                (
                    "test",
                    f"{configuration_id}_test_{metric}",
                    test_results[configuration_id][metric],
                    f"Held-out {metric.replace('_', ' ')} for {configuration_id}.",
                )
            )
    rows.append(
        (
            "method",
            "model_scope",
            "full-feature Random Forest benchmark",
            "Direct full-vocabulary counterpart to proportional Model A.",
        )
    )
    return pd.DataFrame(rows, columns=["section", "item", "value", "description"])


def artifact_descriptions_table() -> pd.DataFrame:
    """Return the six Model B artifact descriptions."""
    files = {
        "fitted_model": MODEL_FILENAME,
        "cross_validation_results": CV_RESULTS_FILENAME,
        "selected_features": FEATURES_FILENAME,
        "test_predictions": TEST_PREDICTIONS_FILENAME,
        "confusion_matrix": CONFUSION_MATRIX_FILENAME,
        "performance_report": PERFORMANCE_REPORT_FILENAME,
    }
    return pd.DataFrame(
        [
            {"filename": files[role], "role": role, "description": ARTIFACT_DESCRIPTIONS[role]}
            for role in files
        ]
    )


def write_performance_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    """Write the same eight-sheet report structure used by Model A."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = MODEL_NAME
    workbook.properties.subject = "Model B cross-validation and held-out evaluation"
    workbook.properties.creator = "IG protocol project"
    table_names = {
        "Model_Summary": "TblModelBSummary",
        "CV_Comparison": "TblFourModelCV",
        "Test_Metrics": "TblModelBTestMetrics",
        "Per_Class_Metrics": "TblModelBPerClass",
        "Confusion_Matrix": "TblModelBConfusion",
        "Selected_Features": "TblModelBFeatures",
        "Test_Predictions": "TblModelBPredictions",
        "Artifact_Descriptions": "TblModelBArtifacts",
    }
    for title, data in sheets.items():
        BASE.add_dataframe_sheet(workbook, title, data, table_names[title])
    workbook.save(path)


def write_outputs(
    *,
    output_dir: Path,
    model_bundle: dict[str, Any],
    cv_results: pd.DataFrame,
    all_features: pd.DataFrame,
    predictions: pd.DataFrame,
    confusion: pd.DataFrame,
    workbook_sheets: dict[str, pd.DataFrame],
) -> dict[str, Path]:
    """Write all six Model B artifacts to a staging directory."""
    paths = {
        "fitted_model": output_dir / MODEL_FILENAME,
        "cross_validation_results": output_dir / CV_RESULTS_FILENAME,
        "selected_features": output_dir / FEATURES_FILENAME,
        "test_predictions": output_dir / TEST_PREDICTIONS_FILENAME,
        "confusion_matrix": output_dir / CONFUSION_MATRIX_FILENAME,
        "performance_report": output_dir / PERFORMANCE_REPORT_FILENAME,
    }
    joblib.dump(model_bundle, paths["fitted_model"], compress=3)
    cv_results.to_csv(paths["cross_validation_results"], index=False, lineterminator="\n")
    all_features.to_csv(paths["selected_features"], index=False, lineterminator="\n")
    predictions.to_csv(paths["test_predictions"], index=False, lineterminator="\n")
    confusion.to_csv(paths["confusion_matrix"], index=False, lineterminator="\n")
    write_performance_workbook(paths["performance_report"], workbook_sheets)
    return paths


def validate_outputs(
    *,
    paths: dict[str, Path],
    test_matrix: sparse.csr_matrix,
    expected_predictions: pd.DataFrame,
    expected_features: pd.DataFrame,
    expected_confusion: pd.DataFrame,
    expected_sheet_names: list[str],
) -> dict[str, Any]:
    """Read back and validate all Model B artifacts."""
    bundle = joblib.load(paths["fitted_model"])
    loaded_predictions = pd.read_csv(
        paths["test_predictions"], dtype=object, keep_default_na=False
    )
    loaded_features = pd.read_csv(
        paths["selected_features"], dtype=object, keep_default_na=False
    )
    loaded_confusion = pd.read_csv(
        paths["confusion_matrix"], dtype=object, keep_default_na=False
    )
    loaded_cv = pd.read_csv(paths["cross_validation_results"], keep_default_na=False)
    workbook = load_workbook(paths["performance_report"], read_only=True, data_only=True)
    try:
        sheet_names = workbook.sheetnames
    finally:
        workbook.close()

    pipelines = bundle.get("pipelines", {})
    reproduced: dict[str, list[str]] = {}
    expected: dict[str, list[str]] = {}
    for configuration_id in CONFIGURATIONS:
        if configuration_id in pipelines:
            reproduced[configuration_id] = pipelines[configuration_id].predict(
                test_matrix
            ).astype(str).tolist()
        expected[configuration_id] = expected_predictions.loc[
            expected_predictions["configuration_id"].eq(configuration_id),
            "predicted_label",
        ].map(str).tolist()
    confusion_columns = [f"predicted_{label}" for label in CLASS_ORDER]
    checks = {
        "all_six_outputs_exist_and_nonempty": len(paths) == 6
        and all(path.is_file() and path.stat().st_size > 0 for path in paths.values()),
        "bundle_contains_B1_and_B2": set(pipelines) == set(CONFIGURATIONS),
        "bundle_class_order_preserved": tuple(bundle.get("class_order", ())) == CLASS_ORDER,
        "both_reloaded_models_reproduce_predictions": reproduced == expected,
        "both_models_use_every_feature": all(
            int(pipeline.named_steps["classifier"].n_features_in_) == test_matrix.shape[1]
            for pipeline in pipelines.values()
        ),
        "prediction_rows_equal_two_times_test_rows": len(loaded_predictions)
        == len(CONFIGURATIONS) * test_matrix.shape[0],
        "prediction_configurations_complete": set(
            loaded_predictions["configuration_id"].map(str)
        )
        == set(CONFIGURATIONS),
        "prediction_labels_approved": set(
            loaded_predictions["predicted_label"].map(str)
        ).issubset(CLASS_ORDER),
        "all_feature_rows_preserved": len(loaded_features) == test_matrix.shape[1],
        "all_feature_indices_sequential": loaded_features[
            "original_tfidf_feature_index"
        ].astype(int).tolist()
        == list(range(test_matrix.shape[1])),
        "all_feature_terms_match": loaded_features["ngram_text"].tolist()
        == expected_features["ngram_text"].map(str).tolist(),
        "confusion_matrices_match_memory": np.array_equal(
            loaded_confusion[confusion_columns].astype(int).to_numpy(),
            expected_confusion[confusion_columns].astype(int).to_numpy(),
        ),
        "each_confusion_total_matches_test_rows": all(
            int(
                loaded_confusion.loc[
                    loaded_confusion["configuration_id"].eq(configuration_id),
                    confusion_columns,
                ].astype(int).to_numpy().sum()
            )
            == test_matrix.shape[0]
            for configuration_id in CONFIGURATIONS
        ),
        "cv_contains_four_configurations": set(loaded_cv["configuration_id"])
        == {"A1", "A2", "B1", "B2"},
        "cv_contains_twenty_fold_rows": int(
            loaded_cv["record_type"].eq("fold").sum()
        )
        == 4 * CV_SPLITS,
        "cv_contains_eight_aggregate_rows": int(
            loaded_cv["record_type"].isin(["mean", "std"]).sum()
        )
        == 8,
        "workbook_has_expected_sheets_in_order": sheet_names == expected_sheet_names,
    }
    checks = {name: bool(value) for name, value in checks.items()}
    failed = [name for name, passed in checks.items() if not passed]
    if failed:
        raise ValueError("Model B output validation failed: " + ", ".join(failed))
    return {"status": "passed", "checks": checks}


def artifact_rows(
    *,
    model_id: str,
    final_dir: Path,
    staged_paths: dict[str, Path],
    created_at: str,
) -> list[dict[str, Any]]:
    """Build artifact registry rows from validated staged files."""
    return [
        {
            "model_id": model_id,
            "artifact_role": role,
            "file_name": path.name,
            "artifact_path": BASE.project_relative(final_dir / path.name),
            "artifact_hash_sha256": BASE.sha256_file(path),
            "size_bytes": path.stat().st_size,
            "artifact_description": ARTIFACT_DESCRIPTIONS[role],
            "created_at": created_at,
        }
        for role, path in staged_paths.items()
    ]


def run_model_b(
    *,
    database_path: Path,
    output_base_dir: Path,
    vectorization_id: str | None = None,
) -> dict[str, Any]:
    """Validate, train, evaluate, export, publish, and register Model B."""
    created_at = utc_now_iso()
    schema_migration = ensure_general_model_registry_schema(database_path)
    vector_event, vector_artifacts = BASE.load_vectorization_event(
        database_path, vectorization_id
    )
    inputs = BASE.load_and_validate_inputs(vector_event, vector_artifacts)
    train_matrix: sparse.csr_matrix = inputs["train_matrix"]
    test_matrix: sparse.csr_matrix = inputs["test_matrix"]
    train_index: pd.DataFrame = inputs["train_index"]
    test_index: pd.DataFrame = inputs["test_index"]
    features: pd.DataFrame = inputs["features"]
    vocabulary_size = train_matrix.shape[1]
    if vocabulary_size != EXPECTED_VOCABULARY_SIZE:
        raise ValueError(
            f"Model B expects {EXPECTED_VOCABULARY_SIZE:,} features; found {vocabulary_size:,}."
        )
    all_features = build_all_features(features)
    model_a_reference = load_model_a_cv_reference(
        database_path, vector_event["vectorization_id"]
    )

    labels = train_index["reviewer_gold_label"].map(str).to_numpy()
    folds, fold_validation = BASE.build_cv_folds(labels)
    b_cv_results, b_cv_summaries = run_cross_validation(train_matrix, labels, folds)
    selected_configuration = select_configuration(b_cv_summaries)
    b_cv_results.loc[
        b_cv_results["configuration_id"].eq(selected_configuration),
        "selected_configuration",
    ] = True

    model_id = BASE.next_model_id(database_path)
    combined_cv = build_four_model_cv_comparison(
        model_a_reference=model_a_reference,
        model_b_id=model_id,
        model_b_results=b_cv_results,
        model_b_selected=selected_configuration,
    )

    pipelines: dict[str, Pipeline] = {}
    test_results: dict[str, dict[str, Any]] = {}
    metric_tables: list[pd.DataFrame] = []
    per_class_tables: list[pd.DataFrame] = []
    confusion_tables: list[pd.DataFrame] = []
    prediction_tables: list[pd.DataFrame] = []
    for configuration_id, configuration in CONFIGURATIONS.items():
        pipeline = build_pipeline(configuration["class_weight"])
        pipeline.fit(train_matrix, labels)
        if int(pipeline.named_steps["classifier"].n_features_in_) != vocabulary_size:
            raise RuntimeError(f"{configuration_id} final model did not use all features.")
        overall, metrics, classes, matrix_table, predictions = build_evaluation_outputs(
            configuration_id, pipeline, test_matrix, test_index
        )
        pipelines[configuration_id] = pipeline
        test_results[configuration_id] = overall
        metric_tables.append(metrics)
        per_class_tables.append(classes)
        confusion_tables.append(matrix_table)
        prediction_tables.append(predictions)
    test_metrics = pd.concat(metric_tables, ignore_index=True)
    per_class = pd.concat(per_class_tables, ignore_index=True)
    confusion = pd.concat(confusion_tables, ignore_index=True)
    predictions = pd.concat(prediction_tables, ignore_index=True)

    output_base_dir.mkdir(parents=True, exist_ok=True)
    final_dir = output_base_dir / model_id
    staged_dir = Path(tempfile.mkdtemp(prefix=f".{model_id}_", dir=output_base_dir))
    model_bundle = {
        "artifact_type": "paired_model_b_full_feature_bundle",
        "model_id": model_id,
        "script_version": SCRIPT_VERSION,
        "class_order": list(CLASS_ORDER),
        "feature_scope": "all_registered_tfidf_features",
        "feature_count": vocabulary_size,
        "cv_selected_primary_configuration": selected_configuration,
        "held_out_reported_configurations": list(CONFIGURATIONS),
        "test_scores_used_for_selection": False,
        "pipelines": pipelines,
    }
    summary = model_summary_table(
        model_id=model_id,
        vector_event=vector_event,
        model_a_reference=model_a_reference,
        created_at=created_at,
        cv_summaries=b_cv_summaries,
        selected_configuration=selected_configuration,
        test_results=test_results,
        training_records=len(train_index),
        test_records=len(test_index),
        vocabulary_size=vocabulary_size,
    )
    workbook_sheets = {
        "Model_Summary": summary,
        "CV_Comparison": combined_cv,
        "Test_Metrics": test_metrics,
        "Per_Class_Metrics": per_class,
        "Confusion_Matrix": confusion,
        "Selected_Features": all_features,
        "Test_Predictions": predictions,
        "Artifact_Descriptions": artifact_descriptions_table(),
    }
    try:
        staged_paths = write_outputs(
            output_dir=staged_dir,
            model_bundle=model_bundle,
            cv_results=combined_cv,
            all_features=all_features,
            predictions=predictions,
            confusion=confusion,
            workbook_sheets=workbook_sheets,
        )
        output_validation = validate_outputs(
            paths=staged_paths,
            test_matrix=test_matrix,
            expected_predictions=predictions,
            expected_features=all_features,
            expected_confusion=confusion,
            expected_sheet_names=list(workbook_sheets),
        )
        correctness_checks = {
            "all_23464_features_used": len(all_features) == 23_464,
            "no_supervised_feature_selection": all(
                list(pipeline.named_steps) == ["classifier"]
                for pipeline in pipelines.values()
            ),
            "both_configurations_fitted_on_all_training_rows": set(pipelines)
            == set(CONFIGURATIONS),
            "each_configuration_predicts_all_test_rows": all(
                len(
                    predictions.loc[
                        predictions["configuration_id"].eq(configuration_id)
                    ]
                )
                == len(test_index)
                for configuration_id in CONFIGURATIONS
            ),
            "each_per_class_support_equals_test_rows": all(
                int(
                    per_class.loc[
                        per_class["configuration_id"].eq(configuration_id), "support"
                    ].sum()
                )
                == len(test_index)
                for configuration_id in CONFIGURATIONS
            ),
            "test_scores_not_used_for_selection": True,
            "same_materialized_folds_as_model_A": model_a_reference["cv_design"][
                "fold_class_distributions"
            ]
            == fold_validation["fold_class_distributions"],
            "four_model_cv_comparison_complete": set(combined_cv["configuration_id"])
            == {"A1", "A2", "B1", "B2"},
        }
        correctness_checks = {
            name: bool(value) for name, value in correctness_checks.items()
        }
        failed = [name for name, passed in correctness_checks.items() if not passed]
        if failed:
            raise RuntimeError("Model B correctness validation failed: " + ", ".join(failed))
        validation_results = {
            "status": "passed",
            "registry_schema_migration": schema_migration,
            "input_validation": inputs["validation"],
            "cross_validation_fold_validation": fold_validation,
            "output_validation": output_validation,
            "final_correctness_checks": correctness_checks,
        }
        artifacts = artifact_rows(
            model_id=model_id,
            final_dir=final_dir,
            staged_paths=staged_paths,
            created_at=created_at,
        )
        input_artifact_records = [
            {
                "artifact_role": role,
                "file_name": values["file_name"],
                "artifact_path": values["artifact_path"],
                "artifact_hash_sha256": values["artifact_hash_sha256"],
                "size_bytes": int(values["size_bytes"]),
            }
            for role, values in sorted(vector_artifacts.items())
        ]
        configuration_records = {
            configuration_id: {
                **RANDOM_FOREST_PARAMETERS,
                "class_weight": configuration["class_weight"],
                "configuration_label": configuration["label"],
                "feature_count": vocabulary_size,
                "feature_selection": None,
            }
            for configuration_id, configuration in CONFIGURATIONS.items()
        }
        confusion_columns = [f"predicted_{label}" for label in CLASS_ORDER]
        final_metrics_registry = {
            "class_order": list(CLASS_ORDER),
            "held_out_reported_configurations": list(CONFIGURATIONS),
            "cv_selected_primary_configuration": selected_configuration,
            "test_scores_used_for_selection": False,
            "configurations": {
                configuration_id: {
                    "configuration_label": CONFIGURATIONS[configuration_id]["label"],
                    "overall": test_results[configuration_id],
                    "per_class": per_class.loc[
                        per_class["configuration_id"].eq(configuration_id)
                    ].to_dict(orient="records"),
                    "confusion_matrix": {
                        "rows_are_actual": True,
                        "columns_are_predicted": True,
                        "values": confusion.loc[
                            confusion["configuration_id"].eq(configuration_id),
                            confusion_columns,
                        ].astype(int).values.tolist(),
                    },
                }
                for configuration_id in CONFIGURATIONS
            },
        }
        cv_registry = {
            "stored_numeric_precision_decimal_places": STORED_METRIC_DECIMALS,
            "model_a_reference_id": model_a_reference["model_id"],
            "model_b_configurations": b_cv_summaries,
            "model_b_records": b_cv_results.to_dict(orient="records"),
            "four_model_comparison_records": combined_cv.to_dict(orient="records"),
        }
        event_values = {
            "model_id": model_id,
            "vectorization_id": vector_event["vectorization_id"],
            "split_id": vector_event["split_id"],
            "matching_id": vector_event["matching_id"],
            "segmentation_id": vector_event["segmentation_id"],
            "document_id": vector_event["document_id"],
            "model_date": created_at[:10],
            "created_at": created_at,
            "model_name": MODEL_NAME,
            "model_designation": MODEL_DESIGNATION,
            "model_description": MODEL_DESCRIPTION,
            "script_path": BASE.project_relative(Path(__file__)),
            "script_version": SCRIPT_VERSION,
            "model_tool": "scikit-learn Pipeline and RandomForestClassifier",
            "model_tool_version": sklearn.__version__,
            "input_artifacts_json": BASE.json_text(input_artifact_records),
            "input_artifacts_count": len(input_artifact_records),
            "class_order_json": BASE.json_text(list(CLASS_ORDER)),
            "class_distributions_json": BASE.json_text(inputs["class_distributions"]),
            "training_records_count": len(train_index),
            "test_records_count": len(test_index),
            "reference_vocabulary_size": vocabulary_size,
            "reference_selected_feature_count": vocabulary_size,
            "retention_proportion": 1.0,
            "retention_calculation": f"{vocabulary_size} / {vocabulary_size} = 1.0",
            "current_vocabulary_size": vocabulary_size,
            "current_selected_feature_count": vocabulary_size,
            "feature_selection_method": "None; all registered TF-IDF features used",
            "feature_selection_parameters_json": BASE.json_text(
                {"method": None, "k": "all", "feature_count": vocabulary_size}
            ),
            "cross_validation_design_json": BASE.json_text(
                {
                    "type": "StratifiedKFold",
                    "n_splits": CV_SPLITS,
                    "shuffle": True,
                    "random_state": RANDOM_STATE,
                    "same_folds_as_model_a": True,
                    "model_a_reference_id": model_a_reference["model_id"],
                    "principal_metric": "macro_f1",
                    "held_out_reported_configurations": list(CONFIGURATIONS),
                    "test_scores_used_for_selection": False,
                    "fold_class_distributions": fold_validation[
                        "fold_class_distributions"
                    ],
                }
            ),
            "random_forest_configurations_json": BASE.json_text(configuration_records),
            "cross_validation_results_json": BASE.json_text(cv_registry),
            "configuration_selection_rule": (
                "Within Model B, select the greatest mean five-fold CV macro-F1 at "
                "12 decimal places and select B1 on a tie. Retain that result as the "
                "primary CV reference, report held-out metrics for B1 and B2, and never "
                "use test performance to change the selection."
            ),
            "selected_configuration": selected_configuration,
            "final_test_metrics_json": BASE.json_text(final_metrics_registry),
            "validation_results_json": BASE.json_text(validation_results),
            "output_artifacts_count": len(artifacts),
            "model_status": "Completed",
            "notes": (
                "Full-feature counterpart to Model A. B1 and B2 retain every registered "
                "TF-IDF column and differ only in class weighting. A1/A2 CV results are "
                f"reused from {model_a_reference['model_id']} after verifying identical "
                "vectorization and fold design; B1/B2 CV is computed in this event. Test "
                "scores are descriptive and do not tune or select a configuration. The "
                "legacy reference-vocabulary registry fields record 23,464-of-23,464 "
                "identity retention for Model B because no Wroblewska proportional "
                "selection applies."
            ),
        }
        registration, final_validation = BASE.register_and_publish(
            database_path=database_path,
            staged_dir=staged_dir,
            final_dir=final_dir,
            event_values=event_values,
            artifacts=artifacts,
            validation_results=validation_results,
        )
    finally:
        if staged_dir.exists():
            shutil.rmtree(staged_dir)

    a_summaries = model_a_reference["summaries"]
    print(f"linked_vectorization_id: {vector_event['vectorization_id']}")
    print(f"model_a_cv_reference_id: {model_a_reference['model_id']}")
    print(f"training_records: {len(train_index)}")
    print(f"test_records: {len(test_index)}")
    print(f"feature_scope: all {vocabulary_size:,} TF-IDF features")
    for configuration_id in ("A1", "A2"):
        print(
            f"{configuration_id}_mean_cv_macro_f1: "
            f"{float(a_summaries[configuration_id]['mean_macro_f1']):.6f}"
        )
    for configuration_id in CONFIGURATIONS:
        print(
            f"{configuration_id}_mean_cv_macro_f1: "
            f"{b_cv_summaries[configuration_id]['mean_macro_f1']:.6f}"
        )
    print(f"model_b_cv_selected_primary_configuration: {selected_configuration}")
    for configuration_id in CONFIGURATIONS:
        print(
            f"{configuration_id}_test_metrics: "
            f"macro_f1={test_results[configuration_id]['macro_f1']:.6f}, "
            f"weighted_f1={test_results[configuration_id]['weighted_f1']:.6f}, "
            f"accuracy={test_results[configuration_id]['accuracy']:.6f}"
        )
    print(f"output_directory: {BASE.project_relative(final_dir)}")
    print(f"registered_model_id: {model_id}")
    print("final_validation_status: passed")
    return {
        "model_id": model_id,
        "vectorization_id": vector_event["vectorization_id"],
        "model_a_reference_id": model_a_reference["model_id"],
        "output_directory": BASE.project_relative(final_dir),
        "feature_count": vocabulary_size,
        "four_model_cv": {
            "A1": model_a_reference["summaries"]["A1"],
            "A2": model_a_reference["summaries"]["A2"],
            "B1": b_cv_summaries["B1"],
            "B2": b_cv_summaries["B2"],
        },
        "selected_configuration": selected_configuration,
        "test_metrics": test_results,
        "per_class_metrics": per_class.to_dict(orient="records"),
        "registration": registration,
        "validation": final_validation,
        "artifacts": artifacts,
    }


def parse_arguments() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(
        description="Train and register paired full-feature Random Forest Model B."
    )
    parser.add_argument(
        "--database", default=str(DATABASE_PATH), help="Corpus inventory SQLite database."
    )
    parser.add_argument(
        "--vectorization-id",
        default=None,
        help="Registered vectorization ID. Default: latest completed vectorization.",
    )
    parser.add_argument(
        "--output-dir",
        default=str(DEFAULT_OUTPUT_DIR),
        help="Base directory for append-only MODEL_###### output folders.",
    )
    return parser.parse_args()


def main() -> None:
    """Command-line entry point."""
    args = parse_arguments()
    try:
        run_model_b(
            database_path=BASE.resolve_project_path(args.database),
            output_base_dir=BASE.resolve_project_path(args.output_dir),
            vectorization_id=args.vectorization_id,
        )
    except Exception as error:
        raise SystemExit(f"Model B failed: {error}") from error


if __name__ == "__main__":
    main()
